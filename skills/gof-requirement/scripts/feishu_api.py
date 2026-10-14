#!/usr/bin/env python3
"""Allowlisted Feishu transport. Credentials stay in the official CLI keychain.

No generic write command, v2 values write, implicit retry, or production target override.
The isolated probe is deliberately separate from the progress-table writer.
Online entry points are parked while the user chooses the Feishu browser route.
Pure local hashing, plan construction and receipt validation remain available.
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

import batch

VERSION = "1.0.98"
OFFICIAL_API_ENABLED = False
DISABLED_REASON = "当前阶段已停用飞书官方 API；使用内置浏览器流程，不发起配置、授权或审批"
PROFILE = "gof-requirement"
ORIGIN = "https://open.feishu.cn"
SCOPES = ("wiki:wiki:readonly", "sheets:spreadsheet:readonly",
          "sheets:spreadsheet:read", "sheets:spreadsheet:write_only",
          "sheets:spreadsheet", "drive:export:readonly", "docx:document:readonly")
PALETTE = ["#bacefd", "#a9efe6", "#e5d6ff", "#fed4a4", "#c4e8a0"]
SECRET_KEYS = re.compile(r"authorization|cookie|password|secret|(?:access|refresh|session|id)[_-]?token", re.I)


class OccupiedRefusal(batch.Blocked):
    """An explicit non-empty-target refusal, not a timeout or permission failure."""


def require_online_enabled():
    batch.require(OFFICIAL_API_ENABLED, DISABLED_REASON)


def reject_secrets(value):
    if isinstance(value, dict):
        batch.require(not any(SECRET_KEYS.search(str(k)) for k in value), "记录中禁止包含认证凭据")
        for item in value.values():
            reject_secrets(item)
    elif isinstance(value, list):
        for item in value:
            reject_secrets(item)
    elif isinstance(value, str):
        batch.require(not re.search(r"\bBearer\s+\S+", value, re.I), "记录中禁止包含认证凭据")


def seal(value, key):
    value = copy.deepcopy(value)
    reject_secrets(value)
    value[key] = batch.digest(value)
    return value


def check_seal(value, key):
    unsealed = dict(value)
    token = unsealed.pop(key, None)
    batch.require(token == batch.digest(unsealed), "接口记录摘要不匹配")
    reject_secrets(value)


class Cli:
    def __init__(self, executable=None, runner=None):
        default = Path(os.environ.get("USERPROFILE", str(Path.home()))) / ".codex/tools/lark-cli" / VERSION / "lark-cli.exe"
        self.executable = str(executable or default)
        self.runner = runner or subprocess.run
        self.calls = []

    def run(self, args, stdin=None, write=False, timeout=60):
        require_online_enabled()
        started = time.monotonic()
        env = dict(os.environ, LARKSUITE_CLI_NO_UPDATE_NOTIFIER="1", LARKSUITE_CLI_NO_SKILLS_NOTIFIER="1")
        try:
            reply = self.runner([self.executable, "--profile", PROFILE, *args], input=stdin, capture_output=True,
                                encoding="utf-8", timeout=timeout, env=env, shell=False)
        except subprocess.TimeoutExpired:
            raise batch.Blocked("接口写入结果未知，必须回读，不可重试" if write else "飞书读取超时") from None
        except OSError:
            raise batch.Blocked("找不到固定版本的官方 lark-cli") from None
        self.calls.append({"command": args[:2], "path": args[2] if len(args) > 2 and args[:2] == ["api", "GET"] else None, "write": write,
                           "elapsed_ms": round((time.monotonic() - started) * 1000), "exit_code": reply.returncode})
        try:
            result = json.loads(reply.stdout.strip() or reply.stderr.strip())
        except (ValueError, TypeError):
            raise batch.Blocked("官方 CLI 返回不可解析数据；写操作须回读核对") from None
        # Never echo stderr, tokens, full auth status, or an untrusted error body.
        if reply.returncode != 0 or result.get("ok", True) is False:
            message = str(result.get("error", {}).get("message", ""))
            if result.get("error", {}).get("type") == "config":
                raise batch.Blocked("飞书专用 gof-requirement profile 尚未配置或不可用；需用户完成应用配置及授权")
            if write and re.search(r"non[- ]?empty|occupied|cannot overwrite|非空|已有内容", message, re.I):
                raise OccupiedRefusal("服务端明确拒绝写入非空目标")
            raise batch.Blocked("官方 CLI 调用失败；写操作须回读核对，检查授权和权限")
        if "code" in result:
            batch.require(result["code"] == 0, "飞书接口业务失败；写操作须回读核对")
        data = result.get("data", result)
        if isinstance(data, dict) and "code" in data:
            batch.require(data["code"] == 0, "飞书接口业务失败")
            data = data.get("data", data)
        self._complete(data)
        return data

    @staticmethod
    def _complete(value):
        if isinstance(value, dict):
            batch.require(value.get("ok", True) is not False and value.get("success", True) is not False
                          and not value.get("errors") and not value.get("failed_count"), "接口含失败项；停止并回读")
            batch.require(value.get("status") not in {"partial_success", "failure", "failed", "error"},
                          "接口返回部分成功或失败；停止并回读，不自动补写")
            batch.require(not any(value.get(k) for k in ("truncated", "is_truncated", "has_more", "hasMore")),
                          "接口数据被截断，不能证明安全")
            for child in value.values():
                Cli._complete(child)
        elif isinstance(value, list):
            for child in value:
                Cli._complete(child)

    def doctor(self):
        require_online_enabled()
        reply = self.runner([self.executable, "--version"], capture_output=True, encoding="utf-8", shell=False)
        batch.require(reply.returncode == 0 and re.search(r"(?<![\d.])1\.0\.98(?![\d.])", reply.stdout),
                      "必须使用已检查的 lark-cli 1.0.98")
        identity = self.run(["auth", "status", "--json", "--verify"])
        batch.require(identity.get("brand") == "feishu", "专用 profile 不是飞书身份，不能切换域名调用")
        user = identity.get("identities", {}).get("user", {})
        batch.require(identity.get("identity") == "user" and identity.get("verified") is True,
                      "飞书 API 尚未以用户身份授权；浏览器登录不能替代 API 授权")
        return {"version": VERSION, "logged_in": True, "identity": "user",
                "user_name": user.get("userName", identity.get("userName")), "required_scopes": list(SCOPES)}

    def api_read(self, path, params=None):
        batch.require(path.startswith("/open-apis/") and ".." not in path and "?" not in path, "非法 API 路径")
        args = ["api", "GET", path, "--as", "user", "--json"]
        if params:
            args += ["--params", batch.canonical(params)]
        return self.run(args)

    def sheet(self, command, token, args=(), stdin=None, write=False):
        permitted = {"+revision-get", "+workbook-export", "+cells-set"}
        batch.require(command in permitted and re.fullmatch(r"[A-Za-z0-9]+", token), "未许可的表格命令")
        if command == "+cells-set":
            batch.require(write and "--allow-overwrite=false" in args, "缺少不可覆盖保护")
            batch.require(stdin is not None and "--writes" in args and "-" in args, "只能发送已校验的区域写入")
        return self.run(["sheets", command, "--spreadsheet-token", token, "--as", "user", "--json", *args],
                        stdin=stdin, write=write, timeout=120 if command == "+workbook-export" else 60)


def revision(cli, token):
    data = cli.sheet("+revision-get", token)
    value = data.get("revision")
    batch.require(type(value) is int and value >= 0, "无法读取文档版本")
    return value


def target(cli):
    data = cli.api_read("/open-apis/wiki/v2/spaces/get_node", {"token": batch.NODE, "obj_type": "wiki"})
    node = data.get("node", {})
    batch.require(node.get("node_token") == batch.NODE and node.get("obj_type") == "sheet", "目标不是获准测试节点的电子表格")
    token = node.get("obj_token", "")
    batch.require(re.fullmatch(r"[A-Za-z0-9]+", token), "缺少真实表格 token")
    return token


def sheets(cli, token):
    data = cli.api_read(f"/open-apis/sheets/v3/spreadsheets/{token}/sheets/query")
    found = data.get("sheets")
    batch.require(isinstance(found, list) and found, "无法列出完整工作表")
    batch.require(all(s.get("resource_type") == "sheet" for s in found), "含无法校验的非网格工作表")
    batch.require(len({s["title"] for s in found}) == len(found), "工作表标题不唯一")
    return found


def rule_projection(records, owned_sheet=None, owned_range=None):
    """Normalize union intervals; exclude only the approved append rectangle."""
    buckets = {}
    area = CellRange(owned_range) if owned_range else None
    for entry in records:
        rule = {k: entry[k] for k in ("type", "options", "multiple", "highlight", "colors")}
        signature = batch.canonical(rule)
        for raw in entry["ranges"]:
            rect = CellRange(raw)
            for col in range(rect.min_col, rect.max_col + 1):
                intervals = [(rect.min_row, rect.max_row)]
                if entry["sheet_id"] == owned_sheet and area and area.min_col <= col <= area.max_col:
                    intervals = [(rect.min_row, min(rect.max_row, area.min_row - 1)),
                                 (max(rect.min_row, area.max_row + 1), rect.max_row)]
                key = (entry["sheet_id"], col, signature)
                buckets.setdefault(key, []).extend((lo, hi) for lo, hi in intervals if lo <= hi)
    result = []
    for (sheet_id, col, signature), intervals in sorted(buckets.items()):
        merged = []
        for lo, hi in sorted(intervals):
            if merged and lo <= merged[-1][1] + 1:
                merged[-1][1] = max(hi, merged[-1][1])
            else:
                merged.append([lo, hi])
        if merged:
            result.append([sheet_id, col, signature, merged])
    return result


def read_rules(cli, token, metadata):
    result = []
    for sheet in metadata:
        sid, grid = sheet["sheet_id"], sheet["grid_properties"]
        rows, cols = grid["row_count"], grid["column_count"]
        batch.require(type(rows) is int and type(cols) is int and 0 < rows * cols <= 4_000_000, "规则读取范围无法证明完整")
        for r in range(1, rows + 1, 5000):
            for c in range(1, cols + 1, 100):
                rect = f"{get_column_letter(c)}{r}:{get_column_letter(min(c + 99, cols))}{min(r + 4999, rows)}"
                data = cli.api_read(f"/open-apis/sheets/v2/spreadsheets/{token}/dataValidation",
                                    {"range": f"{sid}!{rect}", "dataValidationType": "list"})
                batch.require(data.get("sheetId") == sid and isinstance(data.get("dataValidations"), list), "下拉规则返回结构不匹配")
                for rule in data["dataValidations"]:
                    opts = rule.get("options", {})
                    values = rule.get("conditionValues")
                    batch.require(rule.get("dataValidationType") == "list" and isinstance(values, list)
                                  and all(isinstance(v, str) for v in values) and len(values) == len(set(values)),
                                  "下拉来源或选项不支持可靠保留")
                    colors = opts.get("colorValueMap", {})
                    batch.require(isinstance(colors, dict) and set(colors) <= set(values), "下拉颜色映射异常")
                    batch.require(type(opts.get("multipleValues", False)) is bool
                                  and type(opts.get("highlightValidData", False)) is bool
                                  and all(isinstance(color, str) and re.fullmatch(r"#[0-9a-fA-F]{6}", color)
                                          for color in colors.values()), "下拉模式/颜色格式无法可靠保留")
                    ranges = []
                    for raw in rule.get("ranges", []):
                        prefix, coord = raw.split("!", 1)
                        batch.require(prefix == sid, "下拉范围跨表或无法解析")
                        ranges.append(str(CellRange(coord)))
                    batch.require(ranges, "下拉范围为空")
                    result.append({"sheet_id": sid, "ranges": ranges, "type": "list", "options": values,
                                   "multiple": opts.get("multipleValues", False),
                                   "highlight": opts.get("highlightValidData", False), "colors": colors})
    return result


def export(cli, token, output):
    output = Path(output).resolve()
    batch.require(output.suffix.lower() == ".xlsx" and not output.exists(), "快照必须为新的 XLSX 文件")
    cli.sheet("+workbook-export", token, ["--file-extension", "xlsx", "--output-path", str(output)])
    batch.require(output.is_file(), "导出任务未落盘，不能继续")
    return batch.Snapshot(output)


def dropdowns(plan, sid, records):
    result = []
    area = CellRange(plan["range"])
    for spec in plan["dropdowns"]:
        col = CellRange(spec["range"]).min_col
        candidates = []
        for rule in records:
            if rule["sheet_id"] != sid:
                continue
            for raw in rule["ranges"]:
                rect = CellRange(raw)
                if rect.min_col <= col <= rect.max_col and rect.min_row < area.min_row:
                    candidates.append((min(rect.max_row, area.min_row - 1), rule))
        batch.require(candidates, f"{spec['field']} 缺少可保留模式/颜色的原规则")
        nearest = max(end for end, _ in candidates)
        choices = {batch.canonical({k: r[k] for k in ("options", "multiple", "highlight", "colors")}): r
                   for end, r in candidates if end == nearest}
        batch.require(len(choices) == 1, "最近下拉规则冲突，需要澄清")
        rule = next(iter(choices.values()))
        options = rule["options"] + [v for v in spec["options"] if v not in rule["options"]]
        colors = [rule["colors"].get(v, PALETTE[i % len(PALETTE)]) for i, v in enumerate(options)]
        # Existing highlighted options must have an actual color, never a guessed replacement.
        batch.require(not rule["highlight"] or all(v in rule["colors"] for v in rule["options"]), "原下拉配色未完整读取")
        result.append({"field": spec["field"], "range": spec["range"], "options": options,
                       "mode": "multiple" if rule["multiple"] else "single", "highlight": rule["highlight"],
                       "colors": colors, "missing": [v for v in options if v not in rule["options"]]})
    return result


def api_plan(cli, snapshot, plan):
    token = target(cli)
    before = revision(cli, token)
    metadata = sheets(cli, token)
    batch.require([s["title"] for s in sorted(metadata, key=lambda s: s["index"])] == snapshot.book.sheetnames,
                  "API 工作表列表与 XLSX 不匹配")
    sheet = next(s for s in metadata if s["title"] == plan["worksheet"])
    records = read_rules(cli, token, metadata)
    batch.require(before == revision(cli, token), "读取期间表格变化，需要新快照及新计划")
    return seal({"schema": 1, "cli_version": VERSION, "target": batch.TARGET,
                 "batch_plan_id": plan["plan_id"], "baseline_digest": snapshot.fingerprint,
                 "spreadsheet_token": token, "sheet_id": sheet["sheet_id"], "revision": before,
                 "metadata": metadata, "rules": records, "dropdowns": dropdowns(plan, sheet["sheet_id"], records)}, "api_plan_id")


def build_writes(plan, context):
    check_seal(context, "api_plan_id")
    batch.require(context["batch_plan_id"] == plan["plan_id"] and context["baseline_digest"] == plan["baseline_digest"],
                  "接口计划与批次不匹配")
    batch.check_target(context["target"])
    area = CellRange(plan["range"])
    batch.require(area.max_row - area.min_row + 1 <= 5000, "单批超出官方写入行数上限，请先拆分计划")
    validation = {CellRange(d["range"]).min_col: d for d in context["dropdowns"]}
    writes = []
    for col in range(area.min_col, area.max_col + 1):
        letter = get_column_letter(col)
        values = [row["values"][f"{letter}{row['row']}"] for row in plan["rows"]]
        if all(v is None for v in values):
            continue  # Never send nulls/clears into unused fields.
        batch.require(all(v is not None for v in values), "写入列有空洞，不允许混合清空操作")
        cells = []
        for value in values:
            cell = {"formula": value} if isinstance(value, str) and value.startswith("=") else {"value": value}
            if col in validation:
                d = validation[col]
                cell["data_validation"] = {"type": "list", "items": d["options"],
                                           "support_multiple_values": d["mode"] == "multiple",
                                           "enable_highlight": d["highlight"], "highlight_colors": d["colors"]}
                if d["mode"] == "multiple":
                    cell.pop("value")
                    cell["multiple_values"] = [{"value": value}]
            cells.append([cell])
        writes.append({"sheet_id": context["sheet_id"], "range": f"{letter}{area.min_row}:{letter}{area.max_row}",
                       "cells": cells, "allow_overwrite": False})
    batch.require(0 < len(writes) <= 100, "追加写入范围数不支持")
    batch.require(sum(len(w["cells"]) for w in writes) <= 50000, "单批超过官方 CLI 单元格预算，请先拆分计划")
    return writes


def rule_proof(path, fresh, records, context, plan):
    sid = context["sheet_id"]
    batch.require(rule_projection(context["rules"], sid, plan["range"]) == rule_projection(records, sid, plan["range"]),
                  "API 检测到历史下拉规则/颜色/模式变化")
    for spec in context["dropdowns"]:
        rect = CellRange(spec["range"])
        for row in range(rect.min_row, rect.max_row + 1):
            matching = [r for r in records if r["sheet_id"] == sid
                        and any(f"{get_column_letter(rect.min_col)}{row}" in CellRange(a) for a in r["ranges"])]
            batch.require(len({batch.canonical({k: r[k] for k in ("options", "multiple", "highlight", "colors")}) for r in matching}) == 1,
                          "新增下拉规则缺失/冲突")
            rule = matching[0]
            batch.require(rule["options"] == spec["options"] and rule["multiple"] == (spec["mode"] == "multiple")
                          and rule["highlight"] == spec["highlight"]
                          and (not spec["highlight"] or [rule["colors"].get(v) for v in rule["options"]] == spec["colors"]),
                          "新增下拉选项/颜色/模式不匹配")
    proof = seal({"transport": "feishu_api", "api_plan_id": context["api_plan_id"],
                  "snapshot_digest": fresh.fingerprint, "rules": records}, "proof_id")
    batch.write_json_new(path, proof)
    return proof


def validate_rule_evidence(path, fresh, evidence):
    """Called by batch.verify; recomputes proof rather than trusting success booleans."""
    path, manifest = batch.load_batch(path)
    batch.require("api-plan.json" in manifest["artifacts"], "缺少不可变接口计划")
    context = batch.read_json(path / "api-plan.json")
    check_seal(context, "api_plan_id")
    batch.require(isinstance(evidence, dict) and evidence.get("transport") == "feishu_api", "接口批次需要 API 规则验收证据")
    proof_path = Path(evidence.get("proof", "")).resolve()
    batch.require(path in proof_path.parents and proof_path.is_file(), "规则验收记录不在本批目录内")
    proof = batch.read_json(proof_path)
    check_seal(proof, "proof_id")
    batch.require(proof["api_plan_id"] == context["api_plan_id"] and proof["snapshot_digest"] == fresh.fingerprint,
                  "规则证据与接口计划/当前 XLSX 不匹配")
    # Recompute all comparisons without writing another proof.
    plan, sid = manifest["plan"], context["sheet_id"]
    batch.require(rule_projection(context["rules"], sid, plan["range"]) == rule_projection(proof["rules"], sid, plan["range"]),
                  "历史下拉规则发生变化")
    for spec in context["dropdowns"]:
        rect = CellRange(spec["range"])
        for row in range(rect.min_row, rect.max_row + 1):
            coord = f"{get_column_letter(rect.min_col)}{row}"
            rules = [r for r in proof["rules"] if r["sheet_id"] == sid and any(coord in CellRange(a) for a in r["ranges"])]
            batch.require(rules and all(r["options"] == spec["options"] and r["multiple"] == (spec["mode"] == "multiple")
                          and r["highlight"] == spec["highlight"] and (not spec["highlight"] or
                          [r["colors"].get(v) for v in r["options"]] == spec["colors"]) for r in rules), "新增规则验收失败")
    return context["dropdowns"]


def acceptance(receipt):
    data = batch.read_json(receipt)
    check_seal(data, "receipt_id")
    batch.require(data.get("cli_version") == VERSION and data.get("identity") == "user"
                  and data.get("schema") == 1 and data.get("probe_workbook")
                  and data.get("passed") == ["occupied_value", "empty_formula", "repeated_occupancy"],
                  "飞书防覆盖行为尚未通过隔离实测，禁止接口追加")
    batch.require(data["probe_workbook"] != data.get("production_token"), "安全探测不能使用进度表")
    expected = [["KEEP_EXISTING", '=IF(1=1,"","")']]
    observations = data.get("observations", [])
    batch.require(len(observations) == 3 and [o.get("case") for o in observations] == data["passed"]
                  and all(o.get("rejection") == "non_empty" and o.get("before") == expected
                          and o.get("after") == expected for o in observations), "安全记录缺少每次拒写及原值/公式回读")
    return data


def execute(cli, path, receipt):
    path, manifest = batch.load_batch(path)
    context = batch.read_json(path / "api-plan.json")
    check_seal(context, "api_plan_id")
    approved = acceptance(receipt)
    cli.doctor()
    token = target(cli)
    batch.require(token == context["spreadsheet_token"] and approved["production_token"] == token, "目标与安全验证记录不匹配")
    attempt = path / f"api-attempt-{len(batch.events(path)) + 1:06d}"
    attempt.mkdir(exist_ok=False)
    fresh = None
    started = False
    try:
        batch.require(revision(cli, token) == context["revision"], "表格版本变化，必须更新计划")
        fresh = export(cli, token, attempt / "preflight.xlsx")
        batch.preflight(path, fresh)
        current_rules = read_rules(cli, token, sheets(cli, token))
        batch.require(rule_projection(current_rules) == rule_projection(context["rules"]), "下拉规则变化，必须更新计划")
        batch.require(revision(cli, token) == context["revision"], "检查期间出现并发变化")
        writes = build_writes(manifest["plan"], context)
        batch.write_json_new(attempt / "writes.json", writes)
        batch.record(path, "append_started", evidence={"transport": "feishu_api", "api_plan_id": context["api_plan_id"],
                                                     "request_digest": batch.digest(writes)})
        started = True
        cli.sheet("+cells-set", token, ["--writes", "-", "--allow-overwrite=false"],
                  stdin=batch.canonical(writes), write=True)
        fresh.close()
        fresh = export(cli, token, attempt / "after.xlsx")
        post_revision = revision(cli, token)
        after_rules = read_rules(cli, token, sheets(cli, token))
        batch.require(post_revision == revision(cli, token), "验收期间出现并发变化")
        rule_proof(attempt / "rule-proof.json", fresh, after_rules, context, manifest["plan"])
        evidence = {"transport": "feishu_api", "proof": str(attempt / "rule-proof.json")}
        result = batch.verify(path, fresh, evidence)
        batch.require(result["status"] == "complete", "追加未完整通过验收；不自动补写或回滚")
        batch.record(path, "append_verified", fresh=fresh, evidence={"transport": "feishu_api", "validations": evidence,
                                                                   "calls": cli.calls, "revision": post_revision})
        return {"status": "append_verified", "batch": str(path), "range": manifest["plan"]["range"], "calls": cli.calls}
    except (batch.Blocked, OSError, ValueError, KeyError):
        if started and not batch.status(path)["blocked"]:
            batch.record(path, "blocked", evidence={"reason": "飞书接口写入/验收失败或结果未知；先回读，不重试", "transport": "feishu_api"})
        raise
    finally:
        if fresh:
            fresh.close()


def probe(cli, receipt):
    """Only after a separately approved online test plan; creates a new isolated workbook."""
    batch.require(not Path(receipt).exists(), "隔离验证记录不能覆盖")
    cli.doctor()
    production = target(cli)
    title = "GOFCN_API_SAFETY_" + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_%f")
    created = cli.run(["sheets", "+workbook-create", "--title", title, "--as", "user", "--json"], write=True)
    token = created.get("spreadsheet_token", created.get("spreadsheet", {}).get("spreadsheet_token"))
    batch.require(token and token != production, "隔离表创建结果未知，禁止继续/重复创建")
    sid = sheets(cli, token)[0]["sheet_id"]

    def write(rect, cells):
        payload = [{"sheet_id": sid, "range": rect, "cells": cells, "allow_overwrite": False}]
        return cli.sheet("+cells-set", token, ["--writes", "-", "--allow-overwrite=false"], stdin=batch.canonical(payload), write=True)

    def read():
        reply = cli.api_read(f"/open-apis/sheets/v2/spreadsheets/{token}/values_batch_get",
                             {"ranges": f"{sid}!A1:B1", "valueRenderOption": "Formula"})
        batch.require(len(reply.get("valueRanges", [])) == 1, "探测回读结构不匹配")
        return reply["valueRanges"][0]["values"]

    write("A1:B1", [[{"value": "KEEP_EXISTING"}, {"formula": '=IF(1=1,"","")'}]])
    expected = [["KEEP_EXISTING", '=IF(1=1,"","")']]
    batch.require(read() == expected, "隔离基线写入/公式不匹配")
    passed, observations = [], []
    for name, rect in [("occupied_value", "A1:A1"), ("empty_formula", "B1:B1"), ("repeated_occupancy", "A1:B1")]:
        cols = CellRange(rect).max_col - CellRange(rect).min_col + 1
        rejected = False
        before_values = read()
        try:
            write(rect, [[{"value": "MUST_NOT_REPLACE"} for _ in range(cols)]])
        except OccupiedRefusal:
            rejected = True
        after_values = read()
        batch.require(rejected and before_values == after_values == expected, "防覆盖探测失败，生产接口追加保持禁用")
        passed.append(name)
        observations.append({"case": name, "rejection": "non_empty", "before": before_values, "after": after_values})
    data = seal({"schema": 1, "cli_version": VERSION, "identity": "user", "probe_workbook": token,
                 "production_token": production, "passed": passed, "created_at": datetime.now(timezone.utc).isoformat(),
                 "observations": observations, "calls": cli.calls}, "receipt_id")
    batch.write_json_new(receipt, data)
    return {"status": "safety_verified", "receipt": str(Path(receipt).resolve()),
            "probe_url": f"https://centurygames.feishu.cn/sheets/{token}", "calls": cli.calls}


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cli")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("doctor")
    exp = sub.add_parser("snapshot")
    exp.add_argument("--output", required=True)
    plan_cmd = sub.add_parser("plan")
    plan_cmd.add_argument("--snapshot", required=True)
    plan_cmd.add_argument("--request", required=True)
    plan_cmd.add_argument("--current-sheet")
    prep = sub.add_parser("prepare")
    prep.add_argument("--plan", required=True)
    prep.add_argument("--snapshot", required=True)
    prep.add_argument("--workspace", required=True)
    run = sub.add_parser("append")
    run.add_argument("--batch", required=True)
    run.add_argument("--receipt", required=True)
    test = sub.add_parser("probe")
    test.add_argument("--receipt", required=True)
    doc = sub.add_parser("document")
    doc.add_argument("--url", required=True)
    args = parser.parse_args(argv)
    cli = Cli(args.cli)
    snap = None
    try:
        require_online_enabled()
        if args.command == "doctor":
            result = cli.doctor()
        elif args.command == "probe":
            result = probe(cli, args.receipt)
        elif args.command == "append":
            result = execute(cli, args.batch, args.receipt)
        else:
            cli.doctor()
            token = target(cli)
            if args.command == "snapshot":
                before = revision(cli, token)
                snap = export(cli, token, args.output)
                batch.require(before == revision(cli, token), "导出期间版本变化，快照不可用于执行")
                result = {"snapshot": str(snap.path.resolve()), "snapshot_digest": snap.fingerprint, "revision": before}
            elif args.command == "document":
                url = urlsplit(args.url)
                batch.require(url.scheme == "https" and url.hostname == "centurygames.feishu.cn" and not url.username
                              and not url.password and url.port in (None, 443) and re.fullmatch(r"/(docx|wiki)/[A-Za-z0-9]+", url.path), "不是本次飞书脚本文档链接")
                doc_token = url.path.rsplit("/", 1)[1]
                if url.path.startswith("/wiki/"):
                    node = cli.api_read("/open-apis/wiki/v2/spaces/get_node", {"token": doc_token, "obj_type": "wiki"})["node"]
                    batch.require(node["obj_type"] == "docx", "脚本知识库节点不是新版文档")
                    doc_token = node["obj_token"]
                result = cli.api_read(f"/open-apis/docx/v1/documents/{doc_token}/raw_content")
            else:
                snap = batch.Snapshot(args.snapshot)
                if args.command == "plan":
                    offline = batch.make_plan(snap, batch.read_json(args.request), args.current_sheet)
                    result = {"batch_plan": offline, "api_plan": api_plan(cli, snap, offline)}
                else:
                    approved = batch.read_json(args.plan)
                    recomputed = api_plan(cli, snap, approved["batch_plan"])
                    batch.require(recomputed == approved["api_plan"], "接口目标/版本/规则与获批计划变化")
                    result = batch.prepare(approved["batch_plan"], snap, args.workspace, api_plan=recomputed)
        print(batch.canonical(result))
        return 0
    except (batch.Blocked, OSError, ValueError, KeyError, TypeError) as exc:
        print(batch.canonical({"status": "blocked", "reason": str(exc)}))
        return 2
    finally:
        if snap:
            snap.close()


if __name__ == "__main__":
    raise SystemExit(main())
