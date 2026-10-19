#!/usr/bin/env python3
"""One-off oracle: builds XLSX fixtures with openpyxl and records what the Python implementation answers for them.

The TypeScript port is tested against dev/feishu-sdk/test/fixtures/{*.xlsx,golden.json}. Run once, with the bundled
interpreter, before the Python scripts are deleted:

    python.exe -I -B dev/tests/make_fixtures.py
"""
from __future__ import annotations

import base64
import copy
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "plugins/requirement-auto/skills/gof-requirement/scripts"
OUT = ROOT / "dev/feishu-sdk/test/fixtures"
sys.path.insert(0, str(SCRIPTS))
import batch  # noqa: E402
import cls_api  # noqa: E402
import feishu_api  # noqa: E402

PY = sys.executable
URL = batch.TARGET
OLD, MAIN = "GOFCN创意进度 26年2月", "GOFCN创意进度 26年3月"
HEADERS = ["素材编号", "提炼词", "素材名称", "一级标签", "二级标签", "创意\n形式", "客户端", "供应商", "创意者（审核人）", "创意表达", "创意批次"]


def formula(r):
    return f'=CONCATENATE("GOFCN",A{r},"_",D{r},"_",E{r},"_",F{r},"_",B{r})'


def dv(ws, options, sqref):
    rule = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    rule.add(sqref)
    ws.add_data_validation(rule)


def header(ws):
    for col, text in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, text)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFDDEEFF")
        cell.alignment = Alignment(wrap_text=True)


def build_base():
    wb = openpyxl.Workbook()
    notes = wb.active
    notes.title = "说明"
    notes["A1"] = "说明页，不是进度表"
    old = wb.create_sheet(OLD)
    header(old)
    rows = [(1001, "爆点A", None, "玩法", "卖点", "视频", None, "供应商甲", "张三"),
            (1002, "爆点B", "GOFCN1002_玩法_卖点_视频_爆点B", "玩法", "卖点", "视频", None, "供应商甲", "张三"),
            (1003, "爆点C", "GOFCN1010_玩法_卖点_视频_爆点C", "玩法", "卖点", "视频", None, "供应商乙", "李四"),
            (1001, "爆点D", None, "画面", "体验", "图文", None, "供应商乙", "李四")]
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            if value is not None:
                old.cell(r, c, value)
        if row[2] is None:
            old.cell(r, 3, formula(r))
    main = wb.create_sheet(MAIN)
    header(main)
    data = [(1004, "老爆点1", "玩法", "卖点", "视频", "供应商甲", "张三"),
            (1005, "老爆点2", "画面", "体验", "图文", "供应商乙", "李四"),
            (1006, "老爆点3", "玩法", "卖点", "视频", "供应商甲", "张三"),
            (1007, "老爆点4", "画面", "体验", "图文", "供应商乙", "李四")]
    thin = Side(style="thin", color="FF999999")
    for r, (number, keyword, p, s, form, supplier, creator) in enumerate(data, 2):
        for c, value in zip((1, 2, 4, 5, 6, 8, 9), (number, keyword, p, s, form, supplier, creator)):
            cell = main.cell(r, c, value)
            cell.border = Border(bottom=thin)
        main.cell(r, 3, formula(r))
        main.cell(r, 1).number_format = "0"
    main["B3"].comment = Comment("历史备注", "作者")
    main["H3"].hyperlink = "https://example.com/supplier"
    main.merge_cells("M1:N1")
    main["M1"] = "标题"
    main.row_dimensions[3].height = 24
    main.column_dimensions["A"].width = 12
    main.column_dimensions["L"].hidden = True
    dv(main, ["玩法", "画面", "系统"], "D2:D500")
    dv(main, ["卖点", "体验"], "E2:E500")
    dv(main, ["视频", "图文"], "F2:F500")
    dv(main, ["供应商甲", "供应商乙"], "H2:H3")
    dv(main, ["供应商甲", "供应商乙"], "H4:H5")
    dv(main, ["张三", "李四"], "I2:I500")
    main.auto_filter.ref = "A1:K5"
    return wb


def save(wb, name):
    path = OUT / name
    wb.save(path)
    return path


def reload(path):
    return openpyxl.load_workbook(path)


def apply_plan(wb, plan, filled=None):
    """What a correct append looks like in a fresh export."""
    ws = wb[plan["worksheet"]]
    for row in plan["rows"]:
        for coord, value in row["values"].items():
            if value is not None and (filled is None or coord in filled):
                ws[coord] = value
    for spec in plan["dropdowns"]:
        dv(ws, spec["options"], spec["range"])
    if ws.auto_filter.ref:
        left, right = ws.auto_filter.ref.split(":")[0], ws.auto_filter.ref.split(":")[1]
        ws.auto_filter.ref = f"{left}:{right[0]}{plan['range'].split(':')[1].lstrip('ABCDEFGHIJK')}"
    return wb


REQUEST = {"groups": [
    {"quantity": 2, "material_name": "GOFCN1004_玩法_卖点_视频_新爆点", "keyword": "新爆点", "creative_form": "视频",
     "supplier": "供应商丙", "reviewer": "张三", "script_url": "https://centurygames.feishu.cn/docx/ABC123def",
     "brief": "测试需求一", "cls": {"direction": "方向一", "task_type": "creative_a", "requirement_name": "需求一"}},
    {"quantity": 1, "material_name": "画面_体验_图文_另一个", "keyword": "另一个", "creative_form": "图文",
     "supplier": "供应商甲", "reviewer": "李四", "script_url": "https://centurygames.feishu.cn/wiki/XYZ789",
     "brief": "测试需求二", "cls": {"direction": "方向一", "task_type": "creative_a", "requirement_name": "需求二",
                                "quotation": "100.50", "currency": "USD"}}]}


def bad_requests():
    base = copy.deepcopy(REQUEST)
    cases = {}

    def variant(name, change):
        value = copy.deepcopy(base)
        change(value)
        cases[name] = value
    variant("unknown-top", lambda r: r.update(target="x"))
    variant("zero-quantity", lambda r: r["groups"][0].update(quantity=0))
    variant("float-quantity", lambda r: r["groups"][0].update(quantity=2.0))
    variant("name-mismatch", lambda r: r["groups"][0].update(material_name="玩法_卖点_图文_新爆点"))
    variant("bad-url", lambda r: r["groups"][0].update(script_url="https://evil.example.com/docx/ABC"))
    variant("formula-text", lambda r: r["groups"][0].update(brief="=1+1", keyword="新爆点"))
    variant("formula-keyword", lambda r: r["groups"][0].update(keyword="=新爆点", material_name="玩法_卖点_视频_=新爆点"))
    variant("client-filled", lambda r: r["groups"][0].update(client="X"))
    variant("tab-in-supplier", lambda r: r["groups"][0].update(supplier="a\tb"))
    variant("unknown-group-field", lambda r: r["groups"][0].update(extra=1))
    variant("no-tags", lambda r: r["groups"][0].update(material_name="单段_视频_新爆点"))
    variant("reviewer-mismatch", lambda r: r["groups"][0].update(creator="王五"))
    variant("unknown-cls", lambda r: r["groups"][0]["cls"].update(foo=1))
    variant("missing-field", lambda r: r["groups"][0].pop("supplier"))
    variant("worksheet-missing", lambda r: r.update(worksheet="不存在的表"))
    variant("worksheet-explicit", lambda r: r.update(worksheet=MAIN))
    variant("comma-label", lambda r: r["groups"][0].update(supplier="供,应商"))
    return cases


def run(script, *args):
    reply = subprocess.run([PY, "-I", "-B", str(SCRIPTS / script), *map(str, args)], capture_output=True, encoding="utf-8")
    line = reply.stdout.strip().splitlines()[-1] if reply.stdout.strip() else ""
    if not line:
        raise SystemExit(f"{script} {args}: no output\n{reply.stderr}")
    return json.loads(line), reply.returncode


def strip_plan(plan):
    plan = copy.deepcopy(plan)
    plan.pop("plan_id", None)
    plan.pop("baseline_digest", None)
    return plan


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "requests").mkdir(parents=True)
    golden = {"plan": {}, "plan_errors": {}, "prepare": {}, "verify": {}, "preflight": {}, "canonical": [], "flow": [], "cls": {}}

    # ---- fixtures
    save(build_base(), "base.xlsx")
    base = OUT / "base.xlsx"
    for name, value in {"basic": REQUEST, **bad_requests()}.items():
        write(OUT / "requests" / f"{name}.json", value)

    def variant(name, change):
        wb = reload(base)
        change(wb)
        save(wb, f"plan-{name}.xlsx")
    variant("protected", lambda wb: setattr(wb[MAIN].protection, "sheet", True))
    variant("merged", lambda wb: wb[MAIN].merge_cells("A6:B7"))
    variant("no-formula", lambda wb: [setattr(wb[MAIN].cell(r, 3), "value", f"GOFCN{1000 + r}_x") for r in range(2, 6)])
    variant("bad-id", lambda wb: wb[MAIN].cell(3, 1, "12a"))
    variant("missing-id-formula", lambda wb: wb[MAIN].cell(4, 1, None))
    variant("header-dup", lambda wb: wb[MAIN].cell(1, 12, "素材编号2"))
    variant("no-header", lambda wb: wb[MAIN].cell(1, 9, "其他"))
    variant("non-list-dv", lambda wb: wb[MAIN].data_validations.dataValidation.__getitem__(3).__setattr__("formula1", "$Z$1:$Z$3"))
    variant("tie-sheets", lambda wb: wb.copy_worksheet(wb[MAIN]).__setattr__("title", "GOFCN创意进度 26年3月 副本"))

    # ---- plan goldens (fixture name -> request name)
    scenarios = [("base.xlsx", "basic", None)] + [("base.xlsx", name, None) for name in bad_requests()] + [
        ("base.xlsx", "basic", MAIN), ("base.xlsx", "basic", OLD), ("base.xlsx", "basic", "不存在")] + [
        (f"plan-{n}.xlsx", "basic", None) for n in
        ("protected", "merged", "no-formula", "bad-id", "missing-id-formula", "header-dup", "no-header", "non-list-dv", "tie-sheets")]
    for fixture, request, current in scenarios:
        args = ["plan", "--snapshot", OUT / fixture, "--observed-url", URL, "--request", OUT / "requests" / f"{request}.json"]
        if current:
            args += ["--current-sheet", current]
        result, _ = run("batch.py", *args)
        key = f"{fixture}|{request}|{current or ''}"
        if result.get("status") == "blocked":
            golden["plan_errors"][key] = result["reason"]
        else:
            golden["plan"][key] = strip_plan(result)

    # ---- approved plan, prepare, appended variants
    plan, _ = run("batch.py", "plan", "--snapshot", base, "--observed-url", URL, "--request", OUT / "requests/basic.json")
    plan_file = Path(tempfile.mkdtemp()) / "plan.json"
    plan_file.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    root = Path(tempfile.mkdtemp())
    prepared, _ = run("batch.py", "prepare", "--snapshot", base, "--observed-url", URL, "--plan", plan_file, "--output-root", root)
    batch_dir = Path(prepared["batch"])
    golden["prepare"] = {
        "range": prepared["range"], "csv": prepared["csv"],
        "files": {p.name: p.read_text(encoding="utf-8", newline="") for p in sorted(batch_dir.iterdir())
                  if p.is_file() and p.name not in {"manifest.json", "baseline.xlsx"}}}
    again, _ = run("batch.py", "prepare", "--snapshot", base, "--observed-url", URL, "--plan", plan_file, "--output-root", root)
    golden["prepare"]["second_prepare"] = again.get("reason")
    other = copy.deepcopy(plan)
    other["first_number"] += 1
    other_file = plan_file.with_name("other.json")
    other_file.write_text(json.dumps(other, ensure_ascii=False), encoding="utf-8")
    stale, _ = run("batch.py", "prepare", "--snapshot", base, "--observed-url", URL, "--plan", other_file, "--output-root", root)
    golden["prepare"]["stale_plan"] = stale.get("reason")

    def complete():
        return apply_plan(reload(base), plan)
    save(complete(), "applied-complete.xlsx")
    first_row = plan["rows"][0]["row"]
    save(apply_plan(reload(base), plan, filled={c for c in plan["rows"][0]["values"]}), "applied-partial.xlsx")
    save(apply_plan(reload(base), plan, filled={f"A{r['row']}" for r in plan["rows"]}), "applied-blank-names.xlsx")

    def tamper(name, change):
        wb = complete()
        change(wb)
        save(wb, f"applied-{name}.xlsx")
    tamper("outside-value", lambda wb: wb[MAIN].__setitem__("B3", "被改了"))
    tamper("outside-style", lambda wb: setattr(wb[MAIN]["B4"], "fill", PatternFill("solid", fgColor="FFFF0000")))
    tamper("outside-comment", lambda wb: setattr(wb[MAIN]["B3"], "comment", Comment("改过", "作者")))
    tamper("outside-sheet", lambda wb: wb[OLD].__setitem__("B2", "改了老表"))
    tamper("hidden-row", lambda wb: setattr(wb[MAIN].row_dimensions[3], "hidden", True))
    tamper("merge-changed", lambda wb: wb[MAIN].merge_cells("P1:Q1"))
    tamper("filter-changed", lambda wb: setattr(wb[MAIN].auto_filter, "ref", "A1:K3"))
    tamper("rule-changed", lambda wb: setattr(wb[MAIN].data_validations.dataValidation[0], "formula1", '"玩法,画面,系统,新增"'))
    tamper("bad-value", lambda wb: wb[MAIN].__setitem__(f"D{first_row}", "不在选项里"))
    tamper("occupied-optional", lambda wb: wb[MAIN].__setitem__(f"J{first_row}", "不应有值"))
    tamper("wrong-number", lambda wb: wb[MAIN].__setitem__(f"A{first_row}", 9))
    tamper("renamed-sheet", lambda wb: setattr(wb[OLD], "title", "GOFCN创意进度 26年1月"))
    tamper("extra-sheet", lambda wb: wb.create_sheet("多出来的表"))
    tamper("protected", lambda wb: setattr(wb[MAIN].protection, "sheet", True))
    tamper("no-dropdown", lambda wb: wb[MAIN].data_validations.dataValidation.__delitem__(slice(None)))

    def snapshot_has_row(name):
        wb = reload(base)
        wb[MAIN][f"A{first_row}"] = 1
        save(wb, name)
    snapshot_has_row("occupied-row.xlsx")

    for fixture in sorted(p.name for p in OUT.glob("applied-*.xlsx")) + ["base.xlsx"]:
        result, code = run("batch.py", "verify", "--snapshot", OUT / fixture, "--observed-url", URL, "--batch", batch_dir)
        golden["verify"][fixture] = result
    for fixture in ("base.xlsx", "occupied-row.xlsx", "applied-complete.xlsx"):
        result, _ = run("batch.py", "preflight", "--snapshot", OUT / fixture, "--observed-url", URL, "--batch", batch_dir)
        golden["preflight"][fixture] = result

    # ---- batch state machine on a copy of the prepared batch (events carry times and digests, so only outcomes are kept)
    flow_root = Path(tempfile.mkdtemp())
    flow_prepared, _ = run("batch.py", "prepare", "--snapshot", base, "--observed-url", URL, "--plan", plan_file, "--output-root", flow_root)
    flow_batch = Path(flow_prepared["batch"])
    evidence_dir = Path(tempfile.mkdtemp())

    def evidence(value):
        path = evidence_dir / f"e{len(list(evidence_dir.iterdir()))}.json"
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def step(name, step_name, group=None, ev=None, snapshot=None):
        args = ["record", "--batch", flow_batch, "--step", step_name]
        if group:
            args += ["--group", group]
        if ev is not None:
            args += ["--evidence", evidence(ev)]
        if snapshot:
            args += ["--snapshot", OUT / snapshot, "--observed-url", URL]
        result, code = run("batch.py", *args)
        golden["flow"].append({"name": name, "step": step_name, "group": group, "result": result})
        return result

    def state(name):
        result, _ = run("batch.py", "status", "--batch", flow_batch)
        golden["flow"].append({"name": name, "status": {k: result[k] for k in ("append_verified", "blocked", "groups")}
                               if "groups" in result else result})

    step("cls before append", "cls_started", "01")
    step("secret in evidence", "append_started", None, {"access_token": "x"})
    step("append_started", "append_started")
    step("append_started twice", "append_started")
    step("verified with partial", "append_verified", None, {}, "applied-partial.xlsx")
    step("verified without snapshot", "append_verified")
    step("verified", "append_verified", None, {}, "applied-complete.xlsx")
    state("after append_verified")
    step("append_verified twice", "append_verified", None, {}, "applied-complete.xlsx")
    step("bad group", "cls_started", "09")
    step("cls_started", "cls_started", "01")
    step("cls_started twice", "cls_started", "01")
    step("draft before attempt", "draft_saved", "01", {"confirmed": True, "unpublished": True, "url": "https://creative.centurygames.cn/requirement/1/edit-outer"})
    step("draft bad url", "draft_saved", "01", {"confirmed": True, "unpublished": True, "url": "https://evil.example.com/x"})
    step("handoff no reason", "handoff", "01", {})
    step("bad requirement number", "requirement_number", "01", {"confirmed": True, "requirement_number": "GOFCN123"})
    step("blocked no reason", "blocked", None, {})
    step("unknown step", "frobnicate", "01", {})
    step("cls_started group 2", "cls_started", "02")
    step("handoff group 2", "handoff", "02", {"reason": "人工接手"})
    step("handoff group 2 twice", "handoff", "02", {"reason": "再次"})
    step("number group 2", "requirement_number", "02", {"confirmed": True, "requirement_number": "GOFCN20260100001"})
    step("blocked", "blocked", None, {"reason": "测试阻塞"})
    state("after blocked")
    step("record while blocked", "handoff", "01", {"reason": "x"})
    step("resumed without confirm", "resumed", None, {"reason": "x"}, "applied-complete.xlsx")
    step("resumed", "resumed", None, {"confirmed": True, "reason": "已核对", "cls_reconciled": True}, "applied-complete.xlsx")
    state("after resumed")
    for action, observation in [
            ("fill", {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "ok"}),
            ("publish", {}),
            ("save_draft", {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "ok", "required_complete": True}),
            ("save_draft", {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "ok"}),
            ("save_draft", {"logged_in": False, "can_create": True, "unpublished": True, "required_complete": True}),
            ("save_draft", {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "unknown", "required_complete": True}),
            ("import_scripts", {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "ok"})]:
        result, _ = run("batch.py", "cls-check", "--batch", flow_batch, "--group", "01", "--action", action, "--observations", evidence(observation))
        golden["flow"].append({"name": f"cls-check {action}", "observation": observation, "result": result})
    result, _ = run("batch.py", "cls-check", "--batch", flow_batch, "--group", "09", "--action", "save_draft", "--ui", evidence({}))
    golden["flow"].append({"name": "cls-check bad group", "result": result})
    message, _ = run("batch.py", "message", "--batch", flow_batch)
    golden["flow"].append({"name": "message", "result": message})
    report_out = flow_batch / "report.json"
    report, _ = run("batch.py", "report", "--batch", flow_batch, "--output", report_out)
    golden["flow"].append({"name": "report saved", "result": {k: v for k, v in report.items() if k != "report"}})
    outside, _ = run("batch.py", "report", "--batch", flow_batch, "--output", evidence_dir / "elsewhere.json")
    golden["flow"].append({"name": "report outside", "result": outside})
    full = json.loads(report_out.read_text(encoding="utf-8"))
    golden["flow"].append({"name": "report body", "result": {k: full[k] for k in ("status", "worksheet", "range", "target", "message")}})
    golden["flow"].append({"name": "report groups", "result": [{k: g[k] for k in ("group", "quantity", "steps")} for g in full["groups"]]})

    # ---- tampering with a prepared batch
    tamper_root = Path(tempfile.mkdtemp())
    tampered, _ = run("batch.py", "prepare", "--snapshot", base, "--observed-url", URL, "--plan", plan_file, "--output-root", tamper_root)
    tdir = Path(tampered["batch"])
    (tdir / "append.tsv").write_text("x", encoding="utf-8")
    result, _ = run("batch.py", "status", "--batch", tdir)
    golden["flow"].append({"name": "tampered artifact", "result": result})

    # ---- canonical / digest samples (the CLS adapter in the browser verifies these byte for byte)
    samples = [{"b": 1, "a": [1, 2, {"z": None, "y": True}]}, {"中": "文", "e": "é", "q": '"\\\n'}, {"price": 100}, {"price": 100.5},
               [], {}, {"nested": {"k": ["a", {"b": ""}]}}, "plain", 7, None, {"emoji": "😀", "bmp": "￮"}]
    for value in samples:
        golden["canonical"].append({"value": value, "canonical": batch.canonical(value), "sha256": batch.digest(value)})

    # ---- CLS
    catalog = {
        "schema": 1, "can_create": True, "can_import": True,
        "account": {"id": 7, "name": "张三", "status": 1, "roles": ["creative"], "game_ids": [3]},
        "games": [{"id": 3, "name": "无尽冬日", "code": "GOFCN"}, {"id": 4, "name": "其他", "code": "OTHER"}],
        "users": [{"id": 11, "name": "丙公司（供应商丙）", "status": 1, "roles": ["outsource"]},
                  {"id": 12, "name": "供应商甲", "status": 1, "roles": ["outsource"]},
                  {"id": 13, "name": "张三", "status": 1, "roles": ["creative"]},
                  {"id": 14, "name": "李四", "status": 1, "roles": ["creative"]}],
        "directions": [{"id": 21, "name": "方向一", "game_id": 3}, {"id": 22, "name": "方向一", "game_id": 4}],
        "types": [{"type": "creative", "key": "creative_a", "value": 5}],
        "currencies": [{"key": "USD", "value": "usd"}, {"key": "CNY", "value": "cny"}],
        "companies": [], "departments": []}
    write(OUT / "catalog.json", catalog)
    preview, _ = run("cls_api.py", "--preview", "--request", OUT / "requests/basic.json", "--catalog", OUT / "catalog.json")
    golden["cls"]["preview"] = preview
    for name, change in {
            "no-create": lambda c: c.update(can_create=False),
            "no-game-permission": lambda c: c["account"].update(game_ids=[4]),
            "unknown-supplier": lambda c: c["users"].pop(1),
            "dup-direction": lambda c: c["directions"].append({"id": 23, "name": "方向一", "game_id": 3}),
            "secret": lambda c: c.update(authorization="x")}.items():
        broken = copy.deepcopy(catalog)
        change(broken)
        path = evidence_dir / f"catalog-{name}.json"
        path.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
        result, _ = run("cls_api.py", "--preview", "--request", OUT / "requests/basic.json", "--catalog", path)
        golden["cls"][f"preview-{name}"] = result
    bad_quote = copy.deepcopy(REQUEST)
    bad_quote["groups"][1]["cls"]["quotation"] = "abc"
    quote_path = evidence_dir / "bad-quote.json"
    quote_path.write_text(json.dumps(bad_quote, ensure_ascii=False), encoding="utf-8")
    result, _ = run("cls_api.py", "--preview", "--request", quote_path, "--catalog", OUT / "catalog.json")
    golden["cls"]["preview-bad-quote"] = result

    # operation for save_draft on group 01 of the flow batch (cls_started was recorded above)
    observations = {"logged_in": True, "can_create": True, "unpublished": True, "save_result": "ok", "required_complete": True}
    op_path = flow_batch / "save-operation.json"
    result, _ = run("cls_api.py", "--batch", flow_batch, "--group", "01", "--action", "save_draft", "--catalog", OUT / "catalog.json",
                    "--observations", evidence(observations), "--output", op_path)
    golden["cls"]["operation_save"] = {k: v for k, v in result.items() if k not in {"operation", "operation_id"}}
    operation = json.loads(op_path.read_text(encoding="utf-8"))
    golden["cls"]["operation_save_body"] = {k: operation[k] for k in ("schema", "action", "group", "account_id", "payload", "native_defaults", "expected")}
    result, _ = run("cls_api.py", "--batch", flow_batch, "--group", "01", "--action", "import_scripts", "--catalog", OUT / "catalog.json",
                    "--observations", evidence(observations), "--output", flow_batch / "import-operation.json")
    golden["cls"]["operation_import_before_draft"] = result
    result, _ = run("cls_api.py", "--batch", flow_batch, "--group", "01", "--action", "save_draft", "--catalog", OUT / "catalog.json",
                    "--observations", evidence(observations), "--output", op_path)
    golden["cls"]["operation_exists"] = result

    def receipt_for(op, **over):
        parent = {"id": 5001, "status": -2, "user_id": op["account_id"], "for_internal": 0, "sub_scripts": [], "number": "GOFCN20260100009"}
        parent.update({k: v for k, v in op["payload"].items() if k != "asset_data"})
        for key, value in list(parent.items()):
            if value == "" and key.endswith("_date"):
                parent[key] = None
        receipt = {"schema": 1, "transport": "cls_api", "status": "verified", "unpublished": True, "parent": parent,
                   "operation_id": op["operation_id"], "action": op["action"],
                   "calls": [{"method": "POST", "path": "/api/workflow/script/save", "http_status": 200, "business_code": 0},
                             {"method": "GET", "path": "/api/workflow/script/5001", "http_status": 200, "business_code": 0}]}
        receipt.update(over)
        return receipt

    def try_receipt(name, receipt, ev_over=None, op_over=None):
        rp = flow_batch / f"receipt-{name}.json"
        rp.write_text(json.dumps(receipt, ensure_ascii=False), encoding="utf-8")
        ev = {"confirmed": True, "unpublished": True, "transport": "cls_api", "draft_id": 5001, "operation": str(op_path),
              "receipt": str(rp), "url": "https://creative.centurygames.cn/requirement/5001/edit-outer"}
        ev.update(ev_over or {})
        result, _ = run("batch.py", "record", "--batch", flow_batch, "--step", "draft_saved", "--group", "01", "--evidence", evidence(ev))
        golden["cls"][f"draft-{name}"] = result
        return result
    attempt, _ = run("batch.py", "record", "--batch", flow_batch, "--step", "draft_saved", "--group", "01",
                     "--evidence", evidence({"confirmed": True, "unpublished": True, "transport": "cls_api", "draft_id": 5001, "operation": str(op_path),
                                             "receipt": str(flow_batch / "receipt-none.json"), "url": "https://creative.centurygames.cn/requirement/5001/edit-outer"}))
    golden["cls"]["draft-missing-receipt"] = attempt
    try_receipt("no-attempt", receipt_for(operation))
    attempted, _ = run("batch.py", "record", "--batch", flow_batch, "--step", "cls_api_attempt", "--group", "01",
                       "--evidence", evidence({"action": "save_draft", "operation_id": operation["operation_id"], "observations": observations}))
    golden["cls"]["attempt"] = attempted
    try_receipt("wrong-operation", receipt_for(operation, operation_id="0" * 64))
    try_receipt("wrong-name", receipt_for({**operation, "payload": {**operation["payload"], "name": "别的"}}))
    try_receipt("scripts-in-draft", receipt_for(operation, parent={**receipt_for(operation)["parent"], "sub_scripts": [{"id": 1}]}))
    try_receipt("bad-call", receipt_for(operation, calls=[{"method": "POST", "path": "/api/workflow/script/publish", "http_status": 200, "business_code": 0}]))
    try_receipt("not-draft", receipt_for(operation, parent={**receipt_for(operation)["parent"], "status": 1}))
    try_receipt("wrong-url", receipt_for(operation), {"url": "https://creative.centurygames.cn/requirement/9/edit-outer"})
    try_receipt("wrong-number", receipt_for(operation), {"requirement_number": "GOFCN20260100001"})
    try_receipt("wrong-price", receipt_for(operation, parent={**receipt_for(operation)["parent"], "price": 3}))
    try_receipt("ok", receipt_for(operation), {"requirement_number": "GOFCN20260100009"})
    state_after, _ = run("batch.py", "status", "--batch", flow_batch)
    golden["cls"]["steps_group_01"] = state_after["groups"]["01"]
    message2, _ = run("batch.py", "message", "--batch", flow_batch)
    golden["cls"]["message_after_number"] = message2

    # a secret-looking record is refused by seal / reject_secrets
    golden["seal"] = {}
    for name, value in {"key": {"access_token": "x"}, "bearer": {"a": "Bearer abc"}, "nested": {"a": [{"Cookie": "c"}]},
                        "ok": {"a": "bearer"}, "session-id-token": {"id_token": 1}, "authz": {"Authorization": "x"}}.items():
        try:
            feishu_api.reject_secrets(value)
            golden["seal"][name] = "ok"
        except batch.Blocked as exc:
            golden["seal"][name] = str(exc)
    write(OUT / "golden.json", golden)
    print("fixtures:", len(list(OUT.glob("*.xlsx"))), "golden keys:", {k: len(v) for k, v in golden.items()})


if __name__ == "__main__":
    main()
