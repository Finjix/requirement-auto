import copy
import csv
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPT_DIR = Path(__file__).resolve().parents[1] / "skills" / "gof-requirement" / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))
import batch

HEADERS = ["57334", "创意批次", "素材编号", "提炼词", "片头引用\n（初版编号）",
           "创意引用\n（海外汉化解说/迭代）\n（初版编号）", "素材名称", "一级标签\n（方向）",
           "二级标签\n（卖点）", "创意表达\n（可多选）", "创意形式\n（制作形式）",
           "客户端", "设计师\n供应商", "创意者\n脚本审核人"]


def fixture(path):
    w = openpyxl.Workbook()
    s = w.active
    s.title = "26年1月-GOFCN"
    s.append(HEADERS)
    for row, number in [(2, 1000), (4, 1001)]:
        values = {"C": number, "D": "关键词", "H": "明星", "I": "旧标签",
                  "K": "混剪", "M": "丁杰", "N": "张天伟"}
        for col, value in values.items():
            s[f"{col}{row}"] = value
        s[f"G{row}"] = f'=CONCATENATE("GOFCN",C{row},"_",H{row},"_",I{row},"_",K{row},"_",D{row},L{row})'
    s["A1"].fill = PatternFill("solid", fgColor="DDEEDD")
    s.column_dimensions["Z"].hidden = True
    for col, options in {"H": "明星,玩法", "I": "旧标签,展示向", "K": "混剪,2D",
                         "M": "丁杰,潘越", "N": "张天伟,雷佳威"}.items():
        v = DataValidation(type="list", formula1=f'"{options}"')
        v.add(f"{col}2:{col}20")
        s.add_data_validation(v)
    older = w.create_sheet("25年1月-GOFCN")
    older.append(HEADERS)
    older["C2"] = 999
    for col, value in {"D": "关键词", "H": "明星", "I": "旧标签", "K": "混剪", "M": "丁杰", "N": "张天伟"}.items():
        older[f"{col}2"] = value
    older["G2"] = '=CONCATENATE("GOFCN",C2,"_",H2,"_",I2,"_",K2,"_",D2,L2)'
    w.save(path)
    w.close()


def request(quantity=60):
    return {"groups": [{"quantity": quantity, "material_name": "多明星04_2min_混剪_关键词",
                         "topic": "多明星04_2min", "primary_tag": "多明星04",
                         "secondary_tag": "2min", "creative_form": "混剪", "keyword": "关键词",
                         "supplier": "丁杰", "reviewer": "张天伟",
                         "script_url": "https://centurygames.feishu.cn/wiki/CurrentScript123",
                         "brief": "许君聪、吴樾、曹云金、林子聪，2min，总计60条。",
                         "cls": {"direction": "2D剪辑", "task_type": "视频", "requirement_name": "多明星混剪"}}]}


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="gof-skill-test-")
        self.root = Path(self.temp.name)
        self.file = self.root / "baseline.xlsx"
        fixture(self.file)
        self.snapshot = batch.Snapshot(self.file)
        self.plan = batch.make_plan(self.snapshot, request())

    def tearDown(self):
        self.snapshot.close()
        self.temp.cleanup()

    def edit(self, edit_fn, name="edited.xlsx"):
        path = self.root / name
        w = openpyxl.load_workbook(self.file)
        edit_fn(w)
        w.save(path)
        w.close()
        return batch.Snapshot(path)

    def prepared(self, plan=None):
        return Path(batch.prepare(plan or self.plan, self.snapshot, self.root)["batch"])

    def applied(self, plan=None, partial=False, validations=True, path="applied.xlsx"):
        plan = plan or self.plan

        def apply(w):
            s = w[plan["worksheet"]]
            for row in plan["rows"][:1] if partial else plan["rows"]:
                for coord, value in row["values"].items():
                    s[coord] = value
            if validations:
                for spec in plan["dropdowns"]:
                    v = DataValidation(type="list", formula1='"' + ",".join(spec["options"]) + '"')
                    v.add(spec["range"])
                    s.add_data_validation(v)
        return self.edit(apply, path)

    def verified(self):
        path = self.prepared()
        batch.record(path, "append_started")
        fresh = self.applied()
        try:
            batch.record(path, "append_verified", fresh=fresh)
        finally:
            fresh.close()
        return path

    def evidence(self):
        path = self.root / "screenshot.png"
        path.write_bytes(b"isolated test evidence; not an online screenshot")
        return {"confirmed": True, "unpublished": True,
                "url": "https://creative.centurygames.cn/requirement/draft/123", "screenshot": str(path)}

    def test_sixty_rows_and_continuous_ids(self):
        self.assertEqual(self.plan["quantity"], 60)
        self.assertEqual(self.plan["range"], "C5:N64")
        self.assertEqual([r["number"] for r in self.plan["rows"]], list(range(1002, 1062)))
        self.assertEqual(self.plan["groups"][0]["names"][0], "GOFCN1002_多明星04_2min_混剪_关键词")
        self.assertIn("总计60条", self.plan["groups"][0]["request"]["brief"])

    def test_csv_bom_crlf_fixed_tag_and_first_link(self):
        path = self.prepared()
        data = (path / self.plan["groups"][0]["csv"]).read_bytes()
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(data.count(b"\r\n"), 61)
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))
        self.assertEqual(rows[0], batch.CSV_HEADERS)
        self.assertEqual(len(rows) - 1, 60)
        self.assertEqual(rows[1][1], request()["groups"][0]["script_url"])
        self.assertTrue(all(r[1] == "同上" for r in rows[2:]))
        self.assertTrue(all(r[2] == batch.FIXED_TAG for r in rows[1:]))

    def test_csv_quotes_and_comma(self):
        req = request(2)
        req["groups"][0]["keyword"] = '关键词,含"引号"'
        req["groups"][0]["material_name"] = '多明星04_2min_混剪_关键词,含"引号"'
        plan = batch.make_plan(self.snapshot, req)
        path = self.prepared(plan)
        with (path / plan["groups"][0]["csv"]).open(encoding="utf-8-sig", newline="") as f:
            self.assertEqual(list(csv.reader(f))[1][0], plan["groups"][0]["names"][0])

    def test_multiple_supplier_and_topic_files(self):
        req = request(2)
        other = copy.deepcopy(req["groups"][0])
        other.update({"quantity": 3, "supplier": "潘越", "topic": "玩法_2min", "primary_tag": "玩法",
                      "material_name": "玩法_2min_混剪_关键词"})
        req["groups"].append(other)
        plan = batch.make_plan(self.snapshot, req)
        path = self.prepared(plan)
        self.assertEqual(plan["quantity"], 5)
        self.assertEqual(len(list(path.glob("*.csv"))), 2)
        for group in plan["groups"]:
            with (path / group["csv"]).open(encoding="utf-8-sig", newline="") as f:
                self.assertEqual(len(list(csv.reader(f))) - 1, group["request"]["quantity"])

    def test_progress_only_six_fields_inferred_tags_and_blank_other_cells(self):
        group = {key: value for key, value in request(2)["groups"][0].items()
                 if key in batch.REQUIRED_TEXT | {"quantity"}}
        plan = batch.make_plan(self.snapshot, {"groups": [group]})
        normalized = plan["groups"][0]["request"]
        self.assertEqual(normalized["primary_tag"], "多明星04")
        self.assertEqual(normalized["secondary_tag"], "2min")
        self.assertEqual(normalized["creator"], "张天伟")
        self.assertEqual(normalized["cls"], {})
        self.assertEqual(plan["groups"][0]["names"], [
            "GOFCN1002_多明星04_2min_混剪_关键词", "GOFCN1003_多明星04_2min_混剪_关键词"])
        for row in plan["rows"]:
            for col in ("E", "F", "J", "L"):
                self.assertIsNone(row["values"][f"{col}{row['row']}"])
        path = self.prepared(plan)
        with (path / plan["groups"][0]["csv"]).open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))[1:]
        self.assertEqual(rows[0][1], group["script_url"])
        self.assertTrue(all(row[1] == "同上" for row in rows[1:]))
        self.assertTrue(all(row[2] == batch.FIXED_TAG for row in rows))
        fresh = self.applied(plan)
        try:
            self.assertEqual(batch.verify(path, fresh)["status"], "complete")
        finally:
            fresh.close()

    def test_name_sample_number_is_replaced_by_existing_sequence(self):
        req = request(2)
        req["groups"][0]["material_name"] = "GOFCN66988_多明星04_2min_混剪_关键词"
        plan = batch.make_plan(self.snapshot, req)
        self.assertEqual(plan["first_number"], 1002)
        for row in plan["rows"]:
            self.assertTrue(row["name"].startswith(f"GOFCN{row['number']}_"))

    def test_script_content_link_cannot_be_empty_or_invalid(self):
        for link in ("", "   ", None, "同上", "https://example.com/script",
                     "http://centurygames.feishu.cn/wiki/CurrentScript123"):
            req = request(1)
            req["groups"][0]["script_url"] = link
            with self.assertRaises(batch.Blocked):
                batch.make_plan(self.snapshot, req)

    def test_each_clarification_field_required_and_other_progress_fields_blank(self):
        for field in batch.REQUIRED_TEXT:
            req = request(1)
            req["groups"][0].pop(field)
            with self.assertRaises(batch.Blocked):
                batch.make_plan(self.snapshot, req)
        for field in batch.OPTIONAL_FIELDS | {"client"}:
            req = request(1)
            req["groups"][0][field] = "旧行值"
            with self.assertRaises(batch.Blocked):
                batch.make_plan(self.snapshot, req)

    def test_name_and_inferred_tags_must_match_confirmed_fields(self):
        for updates in ({"keyword": "不同提炼词"}, {"primary_tag": "无关标签"},
                        {"material_name": "题材_混剪_关键词", "topic": "题材",
                         "primary_tag": "", "secondary_tag": ""}):
            req = request(1)
            req["groups"][0].update(updates)
            with self.assertRaises(batch.Blocked):
                batch.make_plan(self.snapshot, req)

    def test_reviewer_header_maps_to_progress_reviewer(self):
        fresh = self.edit(lambda w: setattr(w.active["N1"], "value", "审核人"))
        try:
            plan = batch.make_plan(fresh, request(1))
            self.assertEqual(plan["rows"][0]["values"]["N5"], "张天伟")
        finally:
            fresh.close()

    def test_refuse_batch_overwrite_and_artifact_tamper(self):
        path = self.prepared()
        with self.assertRaises(FileExistsError):
            self.prepared()
        (path / "append.tsv").write_text("changed", encoding="utf-8")
        with self.assertRaises(batch.Blocked):
            batch.load_batch(path)

    def test_hidden_content_and_blank_formula_move_tail(self):
        for coord, value in [("Z100", "隐藏列重要内容"), ("Z100", '=""'), ("G100", '=""')]:
            fresh = self.edit(lambda w: setattr(w.active[coord], "value", value), coord + ".xlsx")
            try:
                plan = batch.make_plan(fresh, request(1)) if coord != "G100" else None
                if plan:
                    self.assertEqual(plan["rows"][0]["row"], 101)
                else:
                    with self.assertRaises(batch.Blocked):
                        batch.make_plan(fresh, request(1))
            finally:
                fresh.close()

    def test_fresh_occupancy_and_concurrency_rejected(self):
        path = self.prepared()
        fresh = self.edit(lambda w: setattr(w.active["C5"], "value", 9999))
        try:
            with self.assertRaises(batch.Blocked):
                batch.preflight(path, fresh)
            with self.assertRaises(batch.Blocked):
                batch.prepare(self.plan, fresh, self.root)
        finally:
            fresh.close()

    def test_historical_duplicates_reported_without_repair_and_new_ids_avoid_them(self):
        for coord, value in [("C4", 1000), ("G4", "GOFCN1003_不匹配")]:
            fresh = self.edit(lambda w: setattr(w.active[coord], "value", value), coord + "dup.xlsx")
            try:
                plan = batch.make_plan(fresh, request(1))
                self.assertTrue(plan["history_warnings"])
                self.assertEqual(plan["first_number"], 1001 if coord == "C4" else 1004)
                self.assertEqual(fresh.book.active[coord].value, value)
            finally:
                fresh.close()

    def test_merge_intersection_rejected(self):
        fresh = self.edit(lambda w: w.active.merge_cells("C5:D5"))
        try:
            with self.assertRaises(batch.Blocked):
                batch.make_plan(fresh, request(1))
        finally:
            fresh.close()

    def test_sheet_current_then_explicit_then_latest(self):
        self.assertEqual(self.plan["worksheet"], "26年1月-GOFCN")
        self.assertEqual(batch.make_plan(self.snapshot, request(1), "25年1月-GOFCN")["worksheet"], "25年1月-GOFCN")
        req = request(1)
        req["worksheet"] = "26年1月-GOFCN"
        self.assertEqual(batch.make_plan(self.snapshot, req, "25年1月-GOFCN")["worksheet"], req["worksheet"])
        with self.assertRaises(batch.Blocked):
            batch.make_plan(self.snapshot, request(1), "未知表")

    def test_latest_ambiguity_rejected(self):
        fresh = self.edit(lambda w: w.copy_worksheet(w.active).__setattr__("title", "26年1月-GOFCN-另一个"))
        try:
            with self.assertRaises(batch.Blocked):
                batch.make_plan(fresh, request(1))
        finally:
            fresh.close()

    def test_global_max_from_older_sheet(self):
        def edit(w):
            w["25年1月-GOFCN"]["C2"] = 2000
            w["25年1月-GOFCN"]["G2"] = "GOFCN2000_历史"
        fresh = self.edit(edit)
        try:
            self.assertEqual(batch.make_plan(fresh, request(1))["first_number"], 2001)
        finally:
            fresh.close()

    def test_formula_keyword_injection_and_missing_details_rejected(self):
        for mutate in [lambda r: r["groups"][0].pop("supplier"),
                       lambda r: r["groups"][0].update(quantity=True),
                       lambda r: r["groups"][0].update(keyword="=IMPORTXML(1)"),
                       lambda r: r.update(target_url="https://example.com"),
                       lambda r: r["groups"][0].update(client="\tAPP")]:
            req = request(1)
            mutate(req)
            with self.assertRaises(batch.Blocked):
                batch.make_plan(self.snapshot, req)

    def test_target_node_and_csv_snapshot_rejected(self):
        batch.check_target(batch.TARGET + "?sheet=ELbWAS")
        for url in [batch.TARGET.replace(batch.NODE, "Production"), batch.TARGET.replace("https", "http"),
                    batch.TARGET.replace("centurygames.feishu.cn", "centurygames.feishu.cn.evil.com")]:
            with self.assertRaises(batch.Blocked):
                batch.check_target(url)
        with self.assertRaises(batch.Blocked):
            batch.Snapshot(self.root / "snapshot.csv")

    def test_partial_not_applied_complete_and_resume_guards(self):
        path = self.prepared()
        self.assertEqual(batch.verify(path, self.snapshot)["status"], "not_applied")
        self.assertEqual(batch.preflight(path, self.snapshot)["status"], "ready_to_append")
        batch.record(path, "append_started")
        with self.assertRaises(batch.Blocked):
            batch.preflight(path, self.snapshot)
        partial = self.applied(partial=True, path="partial.xlsx")
        full = self.applied(path="full.xlsx")
        try:
            self.assertEqual(batch.verify(path, partial)["status"], "partial")
            with self.assertRaises(batch.Blocked):
                batch.record(path, "append_verified", fresh=partial)
            self.assertEqual(batch.verify(path, full)["status"], "complete")
            batch.record(path, "append_verified", fresh=full)
            self.assertTrue(batch.status(path)["append_verified"])
        finally:
            partial.close()
            full.close()

    def test_old_values_formulas_styles_and_rules_are_protected(self):
        path = self.prepared()
        edits = [lambda w: setattr(w.active["D2"], "value", "被覆盖"),
                 lambda w: setattr(w.active["G2"], "value", "=1"),
                 lambda w: setattr(w.active["A1"], "fill", PatternFill("solid", fgColor="FF0000")),
                 lambda w: setattr(w.active.data_validations.dataValidation[0], "formula1", '"只剩新选项"'),
                 lambda w: setattr(w["25年1月-GOFCN"]["Z9"], "value", "其他表被写入")]
        for i, edit in enumerate(edits):
            fresh = self.edit(edit, f"conflict{i}.xlsx")
            try:
                self.assertEqual(batch.verify(path, fresh)["status"], "conflict")
            finally:
                fresh.close()

    def test_missing_dropdown_after_append_rejected(self):
        path = self.prepared()
        fresh = self.applied(validations=False)
        try:
            self.assertEqual(batch.verify(path, fresh)["status"], "conflict")
        finally:
            fresh.close()

    def test_do_not_lose_old_dropdown_options(self):
        path = self.prepared()
        fresh = self.applied()
        fresh.close()
        w = openpyxl.load_workbook(self.root / "applied.xlsx")
        w.active.data_validations.dataValidation[-len(self.plan["dropdowns"])].formula1 = '"多明星04"'
        w.save(self.root / "lost-options.xlsx")
        w.close()
        fresh = batch.Snapshot(self.root / "lost-options.xlsx")
        try:
            self.assertEqual(batch.verify(path, fresh)["status"], "conflict")
        finally:
            fresh.close()

    def test_cls_login_permission_failure_unknown_import_and_publish(self):
        path = self.verified()
        batch.record(path, "cls_started", "01")
        ui = {"logged_in": True, "can_create": True, "required_complete": True, "unpublished": True}
        self.assertEqual(batch.cls_check(path, "01", "save_draft", ui)["status"], "allowed")
        for changed in [{"logged_in": False}, {"can_create": False}, {"required_complete": False},
                        {"unpublished": False}, {"save_result": "failure"}, {"save_result": "unknown"}]:
            with self.assertRaises(batch.Blocked):
                batch.cls_check(path, "01", "save_draft", {**ui, **changed})
        for action in ["publish", "delete", "confirm_delivery", "upload_video", "update_scripts"]:
            with self.assertRaises(batch.Blocked):
                batch.cls_check(path, "01", action, ui)
        batch.record(path, "draft_saved", "01", self.evidence())
        for extra in [{}, {"draft_saved": True, "unpublished": True, "import_enabled": False, "import_would_publish": False},
                      {"draft_saved": True, "unpublished": True, "import_enabled": True, "import_would_publish": True}]:
            with self.assertRaises(batch.Blocked):
                batch.cls_check(path, "01", "import_scripts", {**ui, **extra})
        allowed = {**ui, "draft_saved": True, "import_enabled": True, "import_would_publish": False}
        self.assertEqual(batch.cls_check(path, "01", "import_scripts", allowed)["status"], "allowed")

    def test_draft_retries_import_counts_and_message_placeholder(self):
        path = self.verified()
        batch.record(path, "cls_started", "01")
        with self.assertRaises(batch.Blocked):
            batch.record(path, "cls_started", "01")
        batch.record(path, "draft_saved", "01", self.evidence())
        with self.assertRaises(batch.Blocked):
            batch.record(path, "draft_saved", "01", self.evidence())
        bad = {**self.evidence(), "quantity": 59, "names": self.plan["groups"][0]["names"]}
        with self.assertRaises(batch.Blocked):
            batch.record(path, "scripts_imported", "01", bad)
        good = {**bad, "quantity": 60}
        batch.record(path, "scripts_imported", "01", good)
        with self.assertRaises(batch.Blocked):
            batch.record(path, "scripts_imported", "01", good)
        message = batch.message(path)["message"]
        self.assertIn("【发布后补填需求编号】", message)
        self.assertIn(f"创意人：{request()['groups'][0]['reviewer']}", message)
        self.assertNotIn("已发布", message)
        batch.record(path, "requirement_number", "01", {"confirmed": True, "requirement_number": "GOFCN20261013001"})
        self.assertIn("GOFCN20261013001", batch.message(path)["message"])

    def test_blocked_event_stops_later_actions_and_event_tamper(self):
        path = self.prepared()
        batch.record(path, "blocked", evidence={"reason": "需要用户处理"})
        with self.assertRaises(batch.Blocked):
            batch.record(path, "append_started")
        event = path / "events" / "000001.json"
        raw = batch.read_json(event)
        raw["evidence"]["reason"] = "修改"
        event.write_text(json.dumps(raw), encoding="utf-8")
        with self.assertRaises(batch.Blocked):
            batch.status(path)

    def test_cli_plan_is_read_only_and_roundtrip_prepare(self):
        req = self.root / "request.json"
        req.write_text(json.dumps(request(2), ensure_ascii=False), encoding="utf-8")
        before = set(self.root.rglob("*"))
        result = subprocess.run([sys.executable, str(SCRIPT_DIR / "batch.py"), "plan", "--request", str(req),
                                 "--snapshot", str(self.file), "--observed-url", batch.TARGET],
                                capture_output=True, encoding="utf-8", env={**__import__('os').environ, "PYTHONIOENCODING": "utf-8"})
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(set(self.root.rglob("*")), before)
        plan = json.loads(result.stdout)
        path = self.prepared(plan)
        tsv = list(csv.reader(io.StringIO((path / "append.tsv").read_text(encoding="utf-8")), delimiter="\t"))
        self.assertEqual(tsv[0][0], "1002")
        self.assertEqual(tsv[0][4], plan["rows"][0]["values"]["G5"])

    def test_reordered_supplier_column_and_json_key_sort_do_not_shift_paste(self):
        def edit(w):
            s = w.active
            for r in [1, 2, 4]:
                s[f"AA{r}"] = s[f"M{r}"].value
                s[f"M{r}"] = None
        fresh = self.edit(edit)
        try:
            plan = batch.make_plan(fresh, request(1))
            # CLI JSON sorts keys; TSV must still use physical column order.
            plan = json.loads(batch.canonical(plan))
            path = Path(batch.prepare(plan, fresh, self.root)["batch"])
            rows = list(csv.reader(io.StringIO((path / "append.tsv").read_text(encoding="utf-8")), delimiter="\t"))
            self.assertEqual(rows[0][0], "1002")
            self.assertEqual(rows[0][-1], "丁杰")
            self.assertEqual(len(rows[0]), 25)
        finally:
            fresh.close()

    def test_baseline_and_approved_plan_tamper_rejected(self):
        altered = copy.deepcopy(self.plan)
        altered["rows"][0]["values"]["C5"] = 1000
        with self.assertRaises(batch.Blocked):
            batch.prepare(altered, self.snapshot, self.root)
        path = self.prepared()
        (path / "baseline.xlsx").write_bytes(b"broken baseline")
        with self.assertRaises(batch.Blocked):
            batch.load_batch(path)

    def test_resume_requires_explicit_decision_and_reconciled_cls(self):
        path = self.verified()
        batch.record(path, "cls_started", "01")
        batch.record(path, "blocked", evidence={"reason": "登录失效"})
        fresh = self.applied()
        try:
            with self.assertRaises(batch.Blocked):
                batch.record(path, "resumed", evidence={"confirmed": True, "reason": "已登录"}, fresh=fresh)
            batch.record(path, "resumed", evidence={"confirmed": True, "reason": "用户要求继续且已登录",
                         "cls_reconciled": True}, fresh=fresh)
            self.assertFalse(batch.status(path)["blocked"])
            self.assertEqual(batch.cls_check(path, "01", "save_draft", {"logged_in": True, "can_create": True,
                             "unpublished": True, "required_complete": True})["status"], "allowed")
        finally:
            fresh.close()

    def test_partial_or_unknown_attempt_cannot_be_resumed(self):
        path = self.prepared()
        batch.record(path, "append_started")
        batch.record(path, "blocked", evidence={"reason": "粘贴不明"})
        partial = self.applied(partial=True)
        try:
            for fresh in [partial, self.snapshot]:
                with self.assertRaises(batch.Blocked):
                    batch.record(path, "resumed", evidence={"confirmed": True, "reason": "试图恢复"}, fresh=fresh)
        finally:
            partial.close()


if __name__ == "__main__":
    unittest.main()
