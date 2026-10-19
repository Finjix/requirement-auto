#!/usr/bin/env python3
"""Allowlisted Feishu transport over the plugin's SDK bridge (server/dist/bridge.mjs, built on the Feishu Node SDK).

The bridge holds the user's login in the user's own directory (LARK_AGENT_USER_DIR) and never prints a token.
No generic write command, v2 values write, implicit retry, or production target override.
The isolated probe is deliberately separate from the progress-table writer.
Pure local hashing, plan construction and receipt validation need no bridge.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

import batch
import skill_runtime

VERSION = skill_runtime.BRIDGE_VERSION
OFFICIAL_API_ENABLED = True
DISABLED_REASON = "飞书接口入口已停用；检查技能包和授权，不切换浏览器操作"
SCOPES = ("wiki:wiki:readonly", "sheets:spreadsheet:readonly",
          "sheets:spreadsheet:read", "sheets:spreadsheet:write_only",
          "sheets:spreadsheet", "drive:export:readonly", "docx:document:readonly")
PALETTE = ["#bacefd", "#a9efe6", "#e5d6ff", "#fed4a4", "#c4e8a0"]
SECRET_KEYS = re.compile(r"authorization|cookie|password|secret|(?:access|refresh|session|id)[_-]?token", re.I)


class OccupiedRefusal(batch.Blocked):
    """An explicit non-empty-target refusal, not a timeout or permission failure."""


def require_online_enabled():
    batch.require(OFFICIAL_API_ENABLED, DISABLED_REASON)


def reject_secrets(value):
    if isinstance(value, dict):
        batch.require(not any(SECRET_KEYS.search(str(k)) for k in value), "记录中禁止包含认证凭据")
        for item in value.values():
            reject_secrets(item)
    elif isinstance(value, list):
        for item in value:
            reject_secrets(item)
    elif isinstance(value, str):
        batch.require(not re.search(r"\bBearer\s+\S+", value, re.I), "记录中禁止包含认证凭据")


def seal(value, key):
    value = copy.deepcopy(value)
    reject_secrets(value)
    value[key] = batch.digest(value)
    return value


def check_seal(value, key):
    unsealed = dict(value)
    token = unsealed.pop(key, None)
    batch.require(token == batch.digest(unsealed), "接口记录摘要不匹配")
    reject_secrets(value)


# Bridge failures whose message is ours (fixed text, already scrubbed of credentials) and may be shown as is.
SHOWN_KINDS = {"auth_required", "config", "identity_mismatch", "scope_missing", "invalid_request", "login_denied", "login_expired"}


class Bridge:
    def __init__(self, command=None, runner=None):
        # An explicit command (a list) and runner are for local transport tests; the entry point always uses node + bridge.mjs.
        self.override = [str(part) for part in command] if command else None
        self.root = Path(__file__).resolve().parents[1]
        self.runtime = None
        self.runner = runner or subprocess.run
        self.calls = []

    def command(self):
        if self.override:
            return self.override
        if self.runtime is None:
            try:
                runtime = skill_runtime.layout(self.root)
                runtime["node"], _ = skill_runtime.find_node()
                self.runtime = runtime
            except (OSError, ValueError, subprocess.SubprocessError) as error:
                raise batch.Blocked(str(error)) from None
        return [str(self.runtime["node"]), str(self.runtime["bridge"])]

    def call(self, name, payload=None, write=False, timeout=60, path=None):
        require_online_enabled()
        started = time.monotonic()
        # Stdin is always given, so a bridge command can never wait on a pipe somebody left open.
        stdin = batch.canonical(payload) if payload is not None else ""
        try:
            reply = self.runner([*self.command(), name], input=stdin, capture_output=True, encoding="utf-8", timeout=timeout,
                                env=skill_runtime.runtime_environment(self.root), shell=False)
        except subprocess.TimeoutExpired:
            raise batch.Blocked("接口写入结果未知，必须回读，不可重试" if write else "飞书读取超时") from None
        except OSError:
            raise batch.Blocked("无法启动插件内的飞书桥；检查 Node.js 24+ 和 server/dist/bridge.mjs") from None
        self.calls.append({"command": [name], "path": path, "write": write,
                           "elapsed_ms": round((time.monotonic() - started) * 1000), "exit_code": reply.returncode})
        # Never echo stderr, tokens, full auth status, or an untrusted error body.
        try:
            result = json.loads(reply.stdout.strip().splitlines()[-1])
        except (ValueError, TypeError, IndexError):
            raise batch.Blocked("飞书桥返回不可解析数据；写操作须回读核对") from None
        batch.require(isinstance(result, dict) and isinstance(result.get("ok"), bool), "飞书桥返回结构不匹配")
        if not result["ok"]:
            error = result.get("error", {})
            error = error if isinstance(error, dict) else {}
            kind, message = error.get("kind"), str(error.get("message", ""))
            if kind == "occupied":
                raise OccupiedRefusal("服务端明确拒绝写入非空目标")
            if kind == "unknown_result" or (write and kind == "network"):
                raise batch.Blocked("接口写入结果未知，必须回读，不可重试")
            if kind in SHOWN_KINDS:
                raise batch.Blocked(f"飞书不可用（{kind}）：{message}")
            raise batch.Blocked("飞书接口调用失败；写操作须回读核对，检查授权和权限")
        data = result.get("data")
        self._complete(data)
        return data

    @staticmethod
    def _complete(value):
        if isinstance(value, dict):
            batch.require(value.get("ok", True) is not False and value.get("success", True) is not False
                          and not value.get("errors") and not value.get("failed_count"), "接口含失败项；停止并回读")
            batch.require(value.get("status") not in {"partial_success", "failure", "failed", "error"},
                          "接口返回部分成功或失败；停止并回读，不自动补写")
            batch.require(not any(value.get(k) for k in ("truncated", "is_truncated", "has_more", "hasMore")),
                          "接口数据被截断，不能证明安全")
            for child in value.values():
                Bridge._complete(child)
        elif isinstance(value, list):
            for child in value:
                Bridge._complete(child)

    def doctor(self):
        require_online_enabled()
        info = self.call("info")
        batch.require(info.get("bridge") == VERSION, "飞书桥版本与技能不匹配")
        batch.require(info.get("app_configured") is True, "未配置飞书应用（FEISHU_APP_ID / FEISHU_APP_SECRET），由宿主或管理员配置")
        status = self.call("auth-status")
        batch.require(status.get("logged_in") is True,
                      "飞书尚未授权：用 feishu_auth_login 工具，或 feishu_api.py login start 取得授权链接交给用户确认")
        self.check_scopes(SCOPES)
        return {"version": VERSION, "logged_in": True, "identity": "user", "user_name": status.get("name"),
                "open_id": status.get("open_id"), "storage": status.get("storage"), "required_scopes": list(SCOPES)}

    def check_scopes(self, required):
        checked = self.call("auth-check-scopes", {"required": list(required)})
        batch.require(checked.get("complete") is True and not checked.get("missing")
                      and set(required) <= set(checked.get("granted", [])), "飞书用户身份缺少本流程必要权限，需重新授权")

    def api_read(self, path, params=None):
        batch.require(path.startswith("/open-apis/") and ".." not in path and "?" not in path, "非法 API 路径")
        return self.call("api-get", {"path": path, "params": params or {}}, path=path)

    def revision(self, token):
        batch.require(re.fullmatch(r"[A-Za-z0-9]+", token), "非法表格 token")
        return self.call("sheets-revision", {"token": token})

    def export(self, token, output):
        batch.require(re.fullmatch(r"[A-Za-z0-9]+", token), "非法表格 token")
        return self.call("sheets-export", {"token": token, "output": str(output)}, timeout=180)

    def cells_set(self, token, writes):
        """The only write: every range goes out with allow_overwrite=false, which the server enforces."""
        batch.require(re.fullmatch(r"[A-Za-z0-9]+", token), "非法表格 token")
        batch.require(isinstance(writes, list) and writes and all(w.get("allow_overwrite") is False for w in writes), "缺少不可覆盖保护")
        return self.call("sheets-cells-set", {"token": token, "writes": writes}, write=True, timeout=120)

    def create_workbook(self, title):
        return self.call("sheets-create", {"title": title}, write=True)


def revision(bridge, token):
    data = bridge.revision(token)
    value = data.get("revision")
    batch.require(type(value) is int and value >= 0, "无法读取文档版本")
    return value


def target(bridge):
    data = bridge.api_read("/open-apis/wiki/v2/spaces/get_node", {"token": batch.NODE, "obj_type": "wiki"})
    node = data.get("node", {})
    batch.require(node.get("node_token") == batch.NODE and node.get("obj_type") == "sheet", "目标不是获准测试节点的电子表格")
    token = node.get("obj_token", "")
    batch.require(re.fullmatch(r"[A-Za-z0-9]+", token), "缺少真实表格 token")
    return token


def sheets(bridge, token):
    data = bridge.api_read(f"/open-apis/sheets/v3/spreadsheets/{token}/sheets/query")
    found = data.get("sheets")
    batch.require(isinstance(found, list) and found, "无法列出完整工作表")
    batch.require(all(s.get("resource_type") == "sheet" for s in found), "含无法校验的非网格工作表")
    batch.require(len({s["title"] for s in found}) == len(found), "工作表标题不唯一")
    return found


def rule_projection(records, owned_sheet=None, owned_range=None):
    """Normalize union intervals; exclude only the approved append rectangle."""
    buckets = {}
    area = CellRange(owned_range) if owned_range else None
    for entry in records:
        rule = {k: entry[k] for k in ("type", "options", "multiple", "highlight", "colors")}
        signature = batch.canonical(rule)
        for raw in entry["ranges"]:
            rect = CellRange(raw)
            for col in range(rect.min_col, rect.max_col + 1):
                intervals = [(rect.min_row, rect.max_row)]
                if entry["sheet_id"] == owned_sheet and area and area.min_col <= col <= area.max_col:
                    intervals = [(rect.min_row, min(rect.max_row, area.min_row - 1)),
                                 (max(rect.min_row, area.max_row + 1), rect.max_row)]
                key = (entry["sheet_id"], col, signature)
                buckets.setdefault(key, []).extend((lo, hi) for lo, hi in intervals if lo <= hi)
    result = []
    for (sheet_id, col, signature), intervals in sorted(buckets.items()):
        merged = []
        for lo, hi in sorted(intervals):
            if merged and lo <= merged[-1][1] + 1:
                merged[-1][1] = max(hi, merged[-1][1])
            else:
                merged.append([lo, hi])
        if merged:
            result.append([sheet_id, col, signature, merged])
    return result


def read_rules(bridge, token, metadata):
    result = []
    for sheet in metadata:
        sid, grid = sheet["sheet_id"], sheet["grid_properties"]
        rows, cols = grid["row_count"], grid["column_count"]
        batch.require(type(rows) is int and type(cols) is int and 0 < rows * cols <= 4_000_000, "规则读取范围无法证明完整")
        for r in range(1, rows + 1, 5000):
            for c in range(1, cols + 1, 100):
                rect = f"{get_column_letter(c)}{r}:{get_column_letter(min(c + 99, cols))}{min(r + 4999, rows)}"
                data = bridge.api_read(f"/open-apis/sheets/v2/spreadsheets/{token}/dataValidation",
                                    {"range": f"{sid}!{rect}", "dataValidationType": "list"})
                batch.require(data.get("sheetId") == sid and isinstance(data.get("dataValidations"), list), "下拉规则返回结构不匹配")
                for rule in data["dataValidations"]:
                    opts = rule.get("options", {})
                    values = rule.get("conditionValues")
                    batch.require(rule.get("dataValidationType") == "list" and isinstance(values, list)
                                  and all(isinstance(v, str) for v in values) and len(values) == len(set(values)),
                                  "下拉来源或选项不支持可靠保留")
                    colors = opts.get("colorValueMap", {})
                    batch.require(isinstance(colors, dict) and set(colors) <= set(values), "下拉颜色映射异常")
                    batch.require(type(opts.get("multipleValues", False)) is bool
                                  and type(opts.get("highlightValidData", False)) is bool
                                  and all(isinstance(color, str) and re.fullmatch(r"#[0-9a-fA-F]{6}", color)
                                          for color in colors.values()), "下拉模式/颜色格式无法可靠保留")
                    ranges = []
                    for raw in rule.get("ranges", []):
                        prefix, coord = raw.split("!", 1)
                        batch.require(prefix == sid, "下拉范围跨表或无法解析")
                        ranges.append(str(CellRange(coord)))
                    batch.require(ranges, "下拉范围为空")
                    result.append({"sheet_id": sid, "ranges": ranges, "type": "list", "options": values,
                                   "multiple": opts.get("multipleValues", False),
                                   "highlight": opts.get("highlightValidData", False), "colors": colors})
    return result


def export(bridge, token, output):
    output = Path(output).resolve()
    batch.require(output.suffix.lower() == ".xlsx" and not output.exists(), "快照必须为新的 XLSX 文件")
    bridge.export(token, output)
    batch.require(output.is_file(), "导出任务未落盘，不能继续；按官方返回的任务引用只读查询，不重新提交导出")
    return batch.Snapshot(output)


def dropdowns(plan, sid, records):
    result = []
    area = CellRange(plan["range"])
    for spec in plan["dropdowns"]:
        col = CellRange(spec["range"]).min_col
        candidates = []
        for rule in records:
            if rule["sheet_id"] != sid:
                continue
            for raw in rule["ranges"]:
                rect = CellRange(raw)
                if rect.min_col <= col <= rect.max_col and rect.min_row < area.min_row:
                    candidates.append((min(rect.max_row, area.min_row - 1), rule))
        batch.require(candidates, f"{spec['field']} 缺少可保留模式/颜色的原规则")
        nearest = max(end for end, _ in candidates)
        choices = {batch.canonical({k: r[k] for k in ("options", "multiple", "highlight", "colors")}): r
                   for end, r in candidates if end == nearest}
        batch.require(len(choices) == 1, "最近下拉规则冲突，需要澄清")
        rule = next(iter(choices.values()))
        options = rule["options"] + [v for v in spec["options"] if v not in rule["options"]]
        colors = [rule["colors"].get(v, PALETTE[i % len(PALETTE)]) for i, v in enumerate(options)]
        # Existing highlighted options must have an actual color, never a guessed replacement.
        batch.require(not rule["highlight"] or all(v in rule["colors"] for v in rule["options"]), "原下拉配色未完整读取")
        result.append({"field": spec["field"], "range": spec["range"], "options": options,
                       "mode": "multiple" if rule["multiple"] else "single", "highlight": rule["highlight"],
                       "colors": colors, "missing": [v for v in options if v not in rule["options"]]})
    return result


def api_plan(bridge, snapshot, plan):
    token = target(bridge)
    before = revision(bridge, token)
    metadata = sheets(bridge, token)
    batch.require([s["title"] for s in sorted(metadata, key=lambda s: s["index"])] == snapshot.book.sheetnames,
                  "API 工作表列表与 XLSX 不匹配")
    sheet = next(s for s in metadata if s["title"] == plan["worksheet"])
    records = read_rules(bridge, token, metadata)
    batch.require(before == revision(bridge, token), "读取期间表格变化，需要新快照及新计划")
    return seal({"schema": 1, "bridge_version": VERSION, "target": batch.TARGET,
                 "batch_plan_id": plan["plan_id"], "baseline_digest": snapshot.fingerprint,
                 "spreadsheet_token": token, "sheet_id": sheet["sheet_id"], "revision": before,
                 "metadata": metadata, "rules": records, "dropdowns": dropdowns(plan, sheet["sheet_id"], records)}, "api_plan_id")


def build_writes(plan, context):
    check_seal(context, "api_plan_id")
    batch.require(context["batch_plan_id"] == plan["plan_id"] and context["baseline_digest"] == plan["baseline_digest"],
                  "接口计划与批次不匹配")
    batch.check_target(context["target"])
    area = CellRange(plan["range"])
    batch.require(area.max_row - area.min_row + 1 <= 5000, "单批超出官方写入行数上限，请先拆分计划")
    validation = {CellRange(d["range"]).min_col: d for d in context["dropdowns"]}
    writes = []
    for col in range(area.min_col, area.max_col + 1):
        letter = get_column_letter(col)
        values = [row["values"][f"{letter}{row['row']}"] for row in plan["rows"]]
        if all(v is None for v in values):
            continue  # Never send nulls/clears into unused fields.
        batch.require(all(v is not None for v in values), "写入列有空洞，不允许混合清空操作")
        cells = []
        for value in values:
            cell = {"formula": value} if isinstance(value, str) and value.startswith("=") else {"value": value}
            if col in validation:
                d = validation[col]
                cell["data_validation"] = {"type": "list", "items": d["options"],
                                           "support_multiple_values": d["mode"] == "multiple",
                                           "enable_highlight": d["highlight"], "highlight_colors": d["colors"]}
                if d["mode"] == "multiple":
                    cell.pop("value")
                    cell["multiple_values"] = [{"value": value}]
            cells.append([cell])
        writes.append({"sheet_id": context["sheet_id"], "range": f"{letter}{area.min_row}:{letter}{area.max_row}",
                       "cells": cells, "allow_overwrite": False})
    batch.require(0 < len(writes) <= 100, "追加写入范围数不支持")
    batch.require(sum(len(w["cells"]) for w in writes) <= 50000, "单批超过单元格预算，请先拆分计划")
    return writes


def rule_proof(path, fresh, records, context, plan):
    sid = context["sheet_id"]
    batch.require(rule_projection(context["rules"], sid, plan["range"]) == rule_projection(records, sid, plan["range"]),
                  "API 检测到历史下拉规则/颜色/模式变化")
    for spec in context["dropdowns"]:
        rect = CellRange(spec["range"])
        for row in range(rect.min_row, rect.max_row + 1):
            matching = [r for r in records if r["sheet_id"] == sid
                        and any(f"{get_column_letter(rect.min_col)}{row}" in CellRange(a) for a in r["ranges"])]
            batch.require(len({batch.canonical({k: r[k] for k in ("options", "multiple", "highlight", "colors")}) for r in matching}) == 1,
                          "新增下拉规则缺失/冲突")
            rule = matching[0]
            batch.require(rule["options"] == spec["options"] and rule["multiple"] == (spec["mode"] == "multiple")
                          and rule["highlight"] == spec["highlight"]
                          and (not spec["highlight"] or [rule["colors"].get(v) for v in rule["options"]] == spec["colors"]),
                          "新增下拉选项/颜色/模式不匹配")
    proof = seal({"transport": "feishu_api", "api_plan_id": context["api_plan_id"],
                  "snapshot_digest": fresh.fingerprint, "rules": records}, "proof_id")
    batch.write_json_new(path, proof)
    return proof


def validate_rule_evidence(path, fresh, evidence):
    """Called by batch.verify; recomputes proof rather than trusting success booleans."""
    path, manifest = batch.load_batch(path)
    batch.require("api-plan.json" in manifest["artifacts"], "缺少不可变接口计划")
    context = batch.read_json(path / "api-plan.json")
    check_seal(context, "api_plan_id")
    batch.require(isinstance(evidence, dict) and evidence.get("transport") == "feishu_api", "接口批次需要 API 规则验收证据")
    proof_path = Path(evidence.get("proof", "")).resolve()
    batch.require(path in proof_path.parents and proof_path.is_file(), "规则验收记录不在本批目录内")
    proof = batch.read_json(proof_path)
    check_seal(proof, "proof_id")
    batch.require(proof["api_plan_id"] == context["api_plan_id"] and proof["snapshot_digest"] == fresh.fingerprint,
                  "规则证据与接口计划/当前 XLSX 不匹配")
    # Recompute all comparisons without writing another proof.
    plan, sid = manifest["plan"], context["sheet_id"]
    batch.require(rule_projection(context["rules"], sid, plan["range"]) == rule_projection(proof["rules"], sid, plan["range"]),
                  "历史下拉规则发生变化")
    for spec in context["dropdowns"]:
        rect = CellRange(spec["range"])
        for row in range(rect.min_row, rect.max_row + 1):
            coord = f"{get_column_letter(rect.min_col)}{row}"
            rules = [r for r in proof["rules"] if r["sheet_id"] == sid and any(coord in CellRange(a) for a in r["ranges"])]
            batch.require(rules and all(r["options"] == spec["options"] and r["multiple"] == (spec["mode"] == "multiple")
                          and r["highlight"] == spec["highlight"] and (not spec["highlight"] or
                          [r["colors"].get(v) for v in r["options"]] == spec["colors"]) for r in rules), "新增规则验收失败")
    return context["dropdowns"]


def acceptance(receipt):
    data = batch.read_json(receipt)
    check_seal(data, "receipt_id")
    batch.require(data.get("bridge_version") == VERSION and data.get("identity") == "user"
                  and data.get("schema") == 2 and data.get("probe_workbook")
                  and data.get("passed") == ["occupied_value", "empty_formula", "repeated_occupancy"],
                  "飞书防覆盖行为尚未通过隔离实测，禁止接口追加")
    batch.require(data["probe_workbook"] != data.get("production_token"), "安全探测不能使用进度表")
    expected = [["KEEP_EXISTING", '=IF(1=1,"","")']]
    observations = data.get("observations", [])
    batch.require(len(observations) == 3 and [o.get("case") for o in observations] == data["passed"]
                  and all(o.get("rejection") == "non_empty" and o.get("before") == expected
                          and o.get("after") == expected for o in observations), "安全记录缺少每次拒写及原值/公式回读")
    return data


def execute(bridge, path, receipt):
    path, manifest = batch.load_batch(path)
    batch.require(not any(event["step"] == "append_started" for event in batch.events(path)),
                  "已有接口写入尝试；只能回读验收，不再次发送")
    context = batch.read_json(path / "api-plan.json")
    check_seal(context, "api_plan_id")
    approved = acceptance(receipt)
    bridge.doctor()
    token = target(bridge)
    batch.require(token == context["spreadsheet_token"] and approved["production_token"] == token, "目标与安全验证记录不匹配")
    attempt = path / f"api-attempt-{len(batch.events(path)) + 1:06d}"
    attempt.mkdir(exist_ok=False)
    fresh = None
    started = False
    try:
        batch.require(revision(bridge, token) == context["revision"], "表格版本变化，必须更新计划")
        fresh = export(bridge, token, attempt / "preflight.xlsx")
        batch.preflight(path, fresh)
        current_rules = read_rules(bridge, token, sheets(bridge, token))
        batch.require(rule_projection(current_rules) == rule_projection(context["rules"]), "下拉规则变化，必须更新计划")
        batch.require(revision(bridge, token) == context["revision"], "检查期间出现并发变化")
        writes = build_writes(manifest["plan"], context)
        batch.write_json_new(attempt / "writes.json", writes)
        batch.record(path, "append_started", evidence={"transport": "feishu_api", "api_plan_id": context["api_plan_id"],
                                                     "request_digest": batch.digest(writes)})
        started = True
        bridge.cells_set(token, writes)
        fresh.close()
        fresh = export(bridge, token, attempt / "after.xlsx")
        post_revision = revision(bridge, token)
        after_rules = read_rules(bridge, token, sheets(bridge, token))
        batch.require(post_revision == revision(bridge, token), "验收期间出现并发变化")
        rule_proof(attempt / "rule-proof.json", fresh, after_rules, context, manifest["plan"])
        evidence = {"transport": "feishu_api", "proof": str(attempt / "rule-proof.json")}
        result = batch.verify(path, fresh, evidence)
        batch.require(result["status"] == "complete", "追加未完整通过验收；不自动补写或回滚")
        batch.record(path, "append_verified", fresh=fresh, evidence={"transport": "feishu_api", "validations": evidence,
                                                                   "calls": bridge.calls, "revision": post_revision,
                                                                   "snapshot": str(fresh.path.resolve()),
                                                                   "snapshot_digest": fresh.fingerprint})
        report_path = attempt / "verification-report.json"
        batch.write_json_new(report_path, batch.report(path))
        return {"status": "append_verified", "batch": str(path), "worksheet": manifest["plan"]["worksheet"],
                "sheet_id": context["sheet_id"], "range": manifest["plan"]["range"],
                "report": str(report_path), "calls": bridge.calls}
    except (batch.Blocked, OSError, ValueError, KeyError):
        if started and not batch.status(path)["blocked"]:
            batch.record(path, "blocked", evidence={"reason": "飞书接口写入/验收失败或结果未知；先回读，不重试", "transport": "feishu_api"})
        raise
    finally:
        if fresh:
            fresh.close()


def probe(bridge, receipt):
    """Only after a separately approved online test plan; creates a new isolated workbook."""
    batch.require(not Path(receipt).exists(), "隔离验证记录不能覆盖")
    bridge.doctor()
    bridge.check_scopes(("sheets:spreadsheet:create",))
    production = target(bridge)
    title = "GOFCN_API_SAFETY_" + datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_%f")
    created = bridge.create_workbook(title)
    token = created.get("spreadsheet_token")
    batch.require(token and token != production, "隔离表创建结果未知，禁止继续/重复创建")
    sid = sheets(bridge, token)[0]["sheet_id"]

    def write(rect, cells):
        payload = [{"sheet_id": sid, "range": rect, "cells": cells, "allow_overwrite": False}]
        return bridge.cells_set(token, payload)

    def read():
        reply = bridge.api_read(f"/open-apis/sheets/v2/spreadsheets/{token}/values_batch_get",
                             {"ranges": f"{sid}!A1:B1", "valueRenderOption": "Formula"})
        batch.require(len(reply.get("valueRanges", [])) == 1, "探测回读结构不匹配")
        return reply["valueRanges"][0]["values"]

    write("A1:B1", [[{"value": "KEEP_EXISTING"}, {"formula": '=IF(1=1,"","")'}]])
    expected = [["KEEP_EXISTING", '=IF(1=1,"","")']]
    batch.require(read() == expected, "隔离基线写入/公式不匹配")
    passed, observations = [], []
    for name, rect in [("occupied_value", "A1:A1"), ("empty_formula", "B1:B1"), ("repeated_occupancy", "A1:B1")]:
        cols = CellRange(rect).max_col - CellRange(rect).min_col + 1
        rejected = False
        before_values = read()
        try:
            write(rect, [[{"value": "MUST_NOT_REPLACE"} for _ in range(cols)]])
        except OccupiedRefusal:
            rejected = True
        after_values = read()
        batch.require(rejected and before_values == after_values == expected, "防覆盖探测失败，生产接口追加保持禁用")
        passed.append(name)
        observations.append({"case": name, "rejection": "non_empty", "before": before_values, "after": after_values})
    data = seal({"schema": 2, "bridge_version": VERSION, "identity": "user", "probe_workbook": token,
                 "production_token": production, "passed": passed, "created_at": datetime.now(timezone.utc).isoformat(),
                 "observations": observations, "calls": bridge.calls}, "receipt_id")
    batch.write_json_new(receipt, data)
    return {"status": "safety_verified", "receipt": str(Path(receipt).resolve()),
            "probe_url": f"https://centurygames.feishu.cn/sheets/{token}", "calls": bridge.calls}


def reconcile(bridge, path):
    """Re-read an existing batch without sending a write, even after an unknown result."""
    path, manifest = batch.load_batch(path)
    context = batch.read_json(path / "api-plan.json")
    check_seal(context, "api_plan_id")
    bridge.doctor()
    token = target(bridge)
    batch.require("cli_version" not in context, "该批次由旧版 lark-cli 传输创建，不能用新传输回读；请用原技能包回读")
    batch.require(token == context["spreadsheet_token"] and context.get("bridge_version") == VERSION, "回读目标或版本与原批次不符")
    output = path / ("api-readback-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ"))
    output.mkdir(exist_ok=False)
    fresh = None
    try:
        before = revision(bridge, token)
        fresh = export(bridge, token, output / "fresh.xlsx")
        rules = read_rules(bridge, token, sheets(bridge, token))
        batch.require(before == revision(bridge, token), "回读期间表格变化，不能生成验收证据")
        rule_proof(output / "rule-proof.json", fresh, rules, context, manifest["plan"])
        evidence = {"transport": "feishu_api", "proof": str(output / "rule-proof.json")}
        result = batch.verify(path, fresh, evidence)
        batch.write_json_new(output / "readback-report.json", {"verification": result, "revision": before,
                                                              "snapshot_digest": fresh.fingerprint, "calls": bridge.calls})
        return {**result, "snapshot": str(fresh.path), "validation_evidence": evidence,
                "report": str(output / "readback-report.json")}
    finally:
        if fresh:
            fresh.close()


def login(bridge, args):
    """The device authorization the MCP tools run, for when those tools are not available. Never prints a token."""
    if args.action == "start":
        return bridge.call("auth-login", {"extra_scopes": ["sheets:spreadsheet:create"] if args.probe_scope else []})
    if args.action == "poll":
        batch.require(0 <= args.wait <= 50, "--wait 必须在 0 到 50 秒之间")
        return bridge.call("auth-poll", {"wait_ms": args.wait * 1000}, timeout=args.wait + 30)
    if args.action == "status":
        return bridge.call("auth-status")
    return bridge.call("auth-logout")


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("doctor")
    exp = sub.add_parser("snapshot")
    exp.add_argument("--output", required=True)
    plan_cmd = sub.add_parser("plan")
    plan_cmd.add_argument("--snapshot", required=True)
    plan_cmd.add_argument("--request", required=True)
    prep = sub.add_parser("prepare")
    prep.add_argument("--plan", required=True)
    prep.add_argument("--snapshot", required=True)
    destination = prep.add_mutually_exclusive_group()
    destination.add_argument("--output-root", help="内部批次根目录；默认是当前用户目录下的 requirement-auto/batches，无用户目录时用系统 temp")
    destination.add_argument("--workspace", help="历史兼容：使用 <工作区>/outputs/gof-requirement")
    run = sub.add_parser("append")
    run.add_argument("--batch", required=True)
    run.add_argument("--receipt", required=True)
    readback = sub.add_parser("reconcile", help="已有批次只读回读，绝不再次写入")
    readback.add_argument("--batch", required=True)
    test = sub.add_parser("probe")
    test.add_argument("--receipt", required=True)
    doc = sub.add_parser("document")
    doc.add_argument("--url", required=True)
    auth = sub.add_parser("login", help="飞书用户授权；feishu_auth_* 工具不可用时的兜底，同一份登录")
    auth.add_argument("action", choices=["start", "poll", "status", "logout"])
    auth.add_argument("--wait", type=int, default=40, help="poll 最长等待秒数（0-50）")
    auth.add_argument("--probe-scope", action="store_true", help="start：另请求 sheets:spreadsheet:create，仅用于已获批的隔离探测")
    args = parser.parse_args(argv)
    bridge = Bridge()
    snap = None
    try:
        require_online_enabled()
        if args.command == "doctor":
            result = bridge.doctor()
        elif args.command == "login":
            result = login(bridge, args)
        elif args.command == "probe":
            result = probe(bridge, args.receipt)
        elif args.command == "append":
            result = execute(bridge, args.batch, args.receipt)
        elif args.command == "reconcile":
            result = reconcile(bridge, args.batch)
        else:
            bridge.doctor()
            token = target(bridge)
            if args.command == "snapshot":
                before = revision(bridge, token)
                snap = export(bridge, token, args.output)
                batch.require(before == revision(bridge, token), "导出期间版本变化，快照不可用于执行")
                result = {"snapshot": str(snap.path.resolve()), "snapshot_digest": snap.fingerprint, "revision": before}
            elif args.command == "document":
                url = urlsplit(args.url)
                batch.require(url.scheme == "https" and url.hostname == "centurygames.feishu.cn" and not url.username
                              and not url.password and url.port in (None, 443) and re.fullmatch(r"/(docx|wiki)/[A-Za-z0-9]+", url.path), "不是本次飞书脚本文档链接")
                doc_token = url.path.rsplit("/", 1)[1]
                if url.path.startswith("/wiki/"):
                    node = bridge.api_read("/open-apis/wiki/v2/spaces/get_node", {"token": doc_token, "obj_type": "wiki"})["node"]
                    batch.require(node["obj_type"] == "docx", "脚本知识库节点不是新版文档")
                    doc_token = node["obj_token"]
                result = bridge.api_read(f"/open-apis/docx/v1/documents/{doc_token}/raw_content")
            else:
                snap = batch.Snapshot(args.snapshot)
                if args.command == "plan":
                    offline = batch.make_plan(snap, batch.read_json(args.request))
                    result = {"batch_plan": offline, "api_plan": api_plan(bridge, snap, offline)}
                else:
                    approved = batch.read_json(args.plan)
                    recomputed = api_plan(bridge, snap, approved["batch_plan"])
                    batch.require(recomputed == approved["api_plan"], "接口目标/版本/规则与获批计划变化")
                    result = batch.prepare(approved["batch_plan"], snap, args.workspace, api_plan=recomputed,
                                           output_root=args.output_root)
        print(batch.canonical(result))
        return 0 if result.get("status") not in {"partial", "conflict", "blocked"} else 2
    except (batch.Blocked, OSError, ValueError, KeyError, TypeError) as exc:
        print(batch.canonical({"status": "blocked", "reason": str(exc)}))
        return 2
    finally:
        if snap:
            snap.close()


if __name__ == "__main__":
    raise SystemExit(main())
